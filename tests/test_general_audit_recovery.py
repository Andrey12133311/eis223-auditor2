"""Temporary SQLite/original-file regressions; no production or network access."""
import asyncio
from contextlib import ExitStack, redirect_stdout
import hashlib
import io
import json
import os
from pathlib import Path
import sqlite3
import sys
import tempfile
import threading
import time
import unittest
from unittest.mock import Mock, patch

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import v332_all_procurements_audit as audit
import v341_precise_legal as legal
from v381_audit_scheduling import notify_evidence

URL = "https://zakupki.gov.ru/223/filestore/public/1.0/download/fz223/file.html?uid="
LATE_REPLY = ("Дата поступления запроса разъяснения: 01.10.2026.\n"
              "Дата публикации ответа в ЕИС: 07.10.2026.")


class App:
    def __init__(self): self.startups = []
    def get(self, path): return lambda handler: handler
    def on_event(self, event):
        def register(handler):
            if event == "startup": self.startups.append(handler)
            return handler
        return register


class DatabaseFixture:
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="general-recovery-")
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)
        self.path = self.folder / "test.sqlite"
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        self.output = io.StringIO()
        self.stack.enter_context(redirect_stdout(self.output))
        self.stack.enter_context(patch.dict(os.environ, {"EIS223_DATA_DIR": str(self.folder)}))
        self.regs = []
        self.originals = {}
        self.contents = {}
        self.ns = {"app": App(), "conn": self.connect,
                   "_verified_set235": lambda: self.regs,
                   "extract_text": lambda path: "\n".join(p.text for p in Document(path).paragraphs),
                   "analyze": lambda p, docs: legal.evaluate(p, docs)[0],
                   "build_checks": lambda p, detail, card, docs: legal.evaluate(p, docs)[1]}
        with self.connect() as con:
            con.execute("PRAGMA journal_mode=WAL")
            con.executescript("""
                CREATE TABLE purchases(reg_number TEXT PRIMARY KEY,customer TEXT,
                    title TEXT,method TEXT,published_at TEXT,submission_end TEXT,
                    last_seen TEXT,documents_checked INTEGER DEFAULT 0,
                    violations INTEGER DEFAULT 0,checks_ok INTEGER DEFAULT 0,
                    checks_attention INTEGER DEFAULT 0);
                CREATE TABLE docs(id INTEGER PRIMARY KEY,reg_number TEXT,name TEXT,
                    url TEXT,doc_type TEXT,published_at TEXT,local_path TEXT,status TEXT,
                    metadata_json TEXT,final_url TEXT,text_chars INTEGER DEFAULT 0);
                CREATE INDEX docs_by_reg ON docs(reg_number,id);
                CREATE TABLE v319_manual_inputs(reg TEXT PRIMARY KEY);
                CREATE TABLE findings(id INTEGER PRIMARY KEY,reg_number TEXT,code TEXT,
                    kind TEXT,title TEXT,law TEXT,law_url TEXT,evidence TEXT,note TEXT,
                    document TEXT,source_url TEXT,location TEXT,match_text TEXT);
                CREATE TABLE checks(id INTEGER PRIMARY KEY,reg_number TEXT,scope TEXT,
                    item TEXT,status TEXT,evidence TEXT,source TEXT,source_url TEXT,
                    document TEXT,field_path TEXT,note TEXT,law TEXT,law_url TEXT,location TEXT);
            """)

    def connect(self):
        con = sqlite3.connect(self.path, check_same_thread=False)
        con.row_factory = sqlite3.Row
        return con

    def purchase(self, number=1, *, status="checked", present=True):
        reg = str(32619000000 + number)
        self.regs.append(reg)
        with self.connect() as con:
            con.execute("INSERT INTO purchases(reg_number,customer,title,method,published_at,submission_end,last_seen) "
                        "VALUES(?,'ООО Тест','Поставка материалов','Открытый конкурс',"
                        "'2026-10-01','2026-10-30','2026-10-19')", (reg,))
        self.document(reg, LATE_REPLY, status=status, present=present)
        return reg

    def document(self, reg, text, *, status="checked", present=True):
        entries = self.originals.setdefault(reg, [])
        uid = f"{len(entries) + 1:032X}"
        document = Document()
        document.add_paragraph(text)
        encoded = io.BytesIO()
        document.save(encoded)
        data = encoded.getvalue()
        digest = hashlib.sha256(data).hexdigest()
        folder = self.folder / "eis223_official_originals" / reg / uid
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / (digest + ".docx")
        self.contents[path] = data
        if present: path.write_bytes(data)
        (folder / "receipt.json").write_text(json.dumps({
            "reg_number": reg, "uid": uid, "extension": ".docx",
            "sha256": digest, "bytes": len(data)}))
        entries.append({"url": URL + uid, "name": f"Разъяснения {uid}.docx"})
        manifest = self.folder / "eis223_public_browser_intake" / reg / "manifest.json"
        manifest.parent.mkdir(parents=True, exist_ok=True)
        manifest.write_text(json.dumps({"reg_number": reg,
            "source": "public_eis_223_browser", "documents": entries}))
        meta = {"source": "official_eis_original", "source_url": URL + uid,
                "uid": uid, "sha256": digest}
        with self.connect() as con:
            con.execute("INSERT INTO docs(reg_number,name,url,doc_type,published_at,local_path,"
                        "status,metadata_json,final_url,text_chars) VALUES(?,?,?,?,?,?,?,?,?,?)",
                        (reg, entries[-1]["name"], URL + uid, "documentation", "2026-10-07",
                         str(path), status, json.dumps(meta), URL + uid, len(text)))
        return path

    def row(self, reg):
        with self.connect() as con:
            if not con.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='v332_audit_queue'").fetchone():
                return None
            result = con.execute("SELECT * FROM v332_audit_queue WHERE campaign=? AND reg=?",
                                 (audit.CAMPAIGN, reg)).fetchone()
        return dict(result) if result and "retry_count" in result.keys() else None

    def findings(self, reg):
        with self.connect() as con:
            return con.execute("SELECT COUNT(*) FROM findings WHERE reg_number=? AND kind='violation'",
                               (reg,)).fetchone()[0]


class QueueRecovery(DatabaseFixture, unittest.TestCase):
    def test_old_locked_failure_is_recovered_without_repeating_current_success(self):
        locked, invalid, completed = [self.purchase(i) for i in (1, 2, 3)]
        audit.initialize(self.ns)
        for reg in (locked, invalid, completed):
            self.assertEqual(audit.claim_next(self.ns), reg)
            audit.mark(self.ns, reg, "completed")
        with self.connect() as con:
            con.execute("UPDATE v332_audit_queue SET state='failed',error='database is locked' WHERE reg=?", (locked,))
            con.execute("UPDATE v332_audit_queue SET state='failed',error='unsupported document' WHERE reg=?", (invalid,))
            before = [tuple(row) for row in con.execute("SELECT * FROM docs ORDER BY id")]
        self.assertEqual(audit.refresh_changed(self.ns), 1)
        self.assertEqual([self.row(reg)["state"] for reg in (locked, invalid, completed)],
                         ["pending", "failed", "completed"])
        self.assertEqual(audit.claim_next(self.ns), locked)
        audit.mark(self.ns, locked, "completed")
        self.assertEqual(audit.refresh_changed(self.ns), 0)
        with self.connect() as con:
            self.assertEqual([tuple(row) for row in con.execute("SELECT * FROM docs ORDER BY id")], before)

    def test_same_id_readiness_and_restored_file_requeue_once(self):
        reg = self.purchase()
        restored = self.document(reg, "Дополнительный оригинал", status="waiting", present=False)
        audit.initialize(self.ns)
        self.assertEqual(audit.claim_next(self.ns), reg)
        audit.mark(self.ns, reg, "completed")
        last_id = self.row(reg)["last_doc_id"]
        with self.connect() as con:
            con.execute("UPDATE docs SET status='downloaded_no_text' WHERE id=?", (last_id,))
        self.assertEqual(audit.refresh_changed(self.ns), 1)
        self.assertEqual(audit.claim_next(self.ns), reg)
        audit.mark(self.ns, reg, "completed")
        self.assertEqual(audit.refresh_changed(self.ns), 0)
        restored.write_bytes(self.contents[restored])
        self.assertEqual(audit.refresh_changed(self.ns), 1)
        self.assertEqual(audit.claim_next(self.ns), reg)
        audit.mark(self.ns, reg, "completed")
        self.assertEqual(self.row(reg)["last_doc_id"], last_id)
        self.assertEqual(audit.refresh_changed(self.ns), 0)

    def test_pre_analysis_snapshot_detects_arrival_during_review(self):
        reg = self.purchase()
        audit.initialize(self.ns)
        self.assertEqual(audit.claim_next(self.ns), reg)
        snapshot = audit.evidence_version(self.ns, reg)
        with self.connect() as con:
            con.execute("UPDATE docs SET text_chars=text_chars+1 WHERE reg_number=?", (reg,))
        audit.mark(self.ns, reg, "completed", {"last_source_version": snapshot})
        self.assertEqual(audit.refresh_changed(self.ns), 1)
        self.assertEqual(audit.claim_next(self.ns), reg)
        audit.mark(self.ns, reg, "completed")
        self.assertEqual(audit.refresh_changed(self.ns), 0)

    def test_source_change_bypasses_bounded_retry_without_hot_loop(self):
        reg = self.purchase()
        audit.initialize(self.ns)
        self.assertEqual(audit.claim_next(self.ns), reg)
        detail = {"last_source_version": audit.evidence_version(self.ns, reg)}
        self.assertEqual(audit.defer_retry(self.ns, reg, detail, "temporary extraction"), 90)
        self.assertIsNone(audit.claim_next(self.ns))
        self.assertEqual(audit.refresh_changed(self.ns), 0)
        with self.connect() as con:
            con.execute("UPDATE docs SET text_chars=text_chars+1 WHERE reg_number=?", (reg,))
        self.assertEqual(audit.refresh_changed(self.ns), 1)
        self.assertEqual(self.row(reg)["next_retry_at"], 0)
        self.assertEqual(audit.claim_next(self.ns), reg)
        with self.connect() as con:
            con.execute("UPDATE v332_audit_queue SET retry_count=100 WHERE reg=?", (reg,))
        self.assertEqual(audit.defer_retry(self.ns, reg, detail, "temporary extraction"), 1800)

    def test_refresh_does_not_hold_write_lock_or_replace_new_reservation(self):
        reg = self.purchase()
        audit.initialize(self.ns)
        self.assertEqual(audit.claim_next(self.ns), reg)
        audit.mark(self.ns, reg, "completed")
        original = audit.source_version
        def concurrent_update(con, number):
            self.assertFalse(con.in_transaction)
            version = original(con, number)
            with self.connect() as writer:
                writer.execute("UPDATE v332_audit_queue SET state='running',updated_at='new-reservation' WHERE reg=?", (reg,))
            return version + "changed"
        with patch.object(audit, "source_version", side_effect=concurrent_update):
            self.assertEqual(audit.refresh_changed(self.ns), 0)
        self.assertEqual(self.row(reg)["state"], "running")

    def test_extended_sqlite_codes_are_retried_without_hiding_other_errors(self):
        for code in (sqlite3.SQLITE_BUSY, sqlite3.SQLITE_LOCKED, 517, 262):
            error = sqlite3.OperationalError("fixture")
            error.sqlite_errorcode = code
            self.assertTrue(audit.transient_sqlite(error))
        self.assertTrue(audit.transient_sqlite(sqlite3.OperationalError("database table is locked")))
        self.assertFalse(audit.transient_sqlite(sqlite3.OperationalError("no such table: docs")))
        self.assertFalse(audit.transient_sqlite(ValueError("database is locked")))


class WorkerRecovery(DatabaseFixture, unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        real_connect = audit.connect
        def bounded_connect(ns):
            con = real_connect(ns)
            con.execute("PRAGMA busy_timeout=50")
            return con
        self.stack.enter_context(patch.object(audit, "connect", side_effect=bounded_connect))
        self.stack.enter_context(patch.object(audit, "RETRY_BASE_SECONDS", .015))
        self.stack.enter_context(patch.object(audit, "RETRY_MAX_SECONDS", .04))
        self.stack.enter_context(patch.object(audit, "QUEUE_POLL_SECONDS", .02))
        self.stack.enter_context(patch.object(audit, "HEARTBEAT_SECONDS", .04))
        self.stack.enter_context(patch.object(audit, "DOCUMENT_RETRY_BASE_SECONDS", .2))
        self.stack.enter_context(patch.object(audit, "DOCUMENT_RETRY_MAX_SECONDS", .4))
        self.real_sleep = asyncio.sleep
        async def fast_start(seconds, *args, **kwargs):
            return await self.real_sleep(0 if seconds == 14 else seconds, *args, **kwargs)
        self.stack.enter_context(patch.object(audit.asyncio, "sleep", side_effect=fast_start))
        self.blockers = []

    async def asyncTearDown(self):
        for blocker in self.blockers:
            blocker.rollback()
            blocker.close()
        task = self.ns.get("_v332_campaign_worker")
        if task is not None:
            task.cancel()
            await asyncio.wait_for(asyncio.gather(task, return_exceptions=True), 5)

    async def until(self, predicate, timeout=5):
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            if predicate(): return
            await self.real_sleep(.005)
        self.fail("Condition timed out. Captured logs:\n" + self.output.getvalue())

    async def start(self):
        audit.install(self.ns)
        await self.ns["app"].startups[0]()
        return self.ns["_v332_campaign_worker"]

    def lock_database(self):
        blocker = self.connect()
        blocker.execute("BEGIN IMMEDIATE")
        self.blockers.append(blocker)
        return blocker

    async def test_real_claim_lock_recovers_all_workers_without_restart_or_new_documents(self):
        regs = [self.purchase(number) for number in (1, 2, 3)]
        original = audit.initialize
        locked = threading.Event()
        def initialized_then_locked(ns):
            result = original(ns)
            self.lock_database()
            locked.set()
            return result
        initialization = self.stack.enter_context(patch.object(audit, "initialize", side_effect=initialized_then_locked))
        campaign = await self.start()
        await self.until(lambda: locked.is_set() and "SQLITE_RETRY phase=claim" in self.output.getvalue())
        self.assertFalse(campaign.done())
        self.assertEqual(audit.runtime_snapshot(self.ns)["live_workers"], 4)
        tasks = dict(self.ns["_v332_tasks"])
        self.assertTrue(all(not task.done() for task in tasks.values()))
        self.assertTrue(all(self.row(reg)["state"] == "pending" for reg in regs))
        with self.connect() as con:
            before = [tuple(row) for row in con.execute("SELECT id,status,metadata_json FROM docs ORDER BY id")]
        self.blockers[0].rollback()
        await self.until(lambda: all(self.row(reg)["state"] == "completed" for reg in regs))
        self.assertEqual(initialization.call_count, 1)
        self.assertIs(campaign, self.ns["_v332_campaign_worker"])
        self.assertTrue(all(self.findings(reg) == 1 for reg in regs))
        self.assertEqual(audit.progress(self.ns)["new_evidence_based_findings"], 3)
        self.assertIn("SQLITE_RECOVERED phase=claim", self.output.getvalue())
        await self.until(lambda: "V332 HEARTBEAT" in self.output.getvalue())
        with self.connect() as con:
            self.assertEqual([tuple(row) for row in con.execute("SELECT id,status,metadata_json FROM docs ORDER BY id")], before)

    async def test_initialization_retries_actual_sqlite_lock(self):
        reg = self.purchase()
        blocker = self.lock_database()
        campaign = await self.start()
        await self.until(lambda: self.ns["_v332_runtime"]["status"] == "retrying_initialization")
        self.assertFalse(campaign.done())
        blocker.rollback()
        await self.until(lambda: self.ns["_v332_runtime"]["status"] == "running" and self.row(reg)["state"] == "completed")
        self.assertEqual(self.findings(reg), 1)

    async def phase_lock(self, phase):
        reg = self.purchase()
        name = "persist" if phase == "persist" else "mark"
        original = getattr(audit, name)
        first = threading.Event()
        analyzer = Mock(wraps=self.ns["analyze"])
        checker = Mock(wraps=self.ns["build_checks"])
        self.stack.enter_context(patch.dict(self.ns, {"analyze": analyzer, "build_checks": checker}))
        def collide(*args, **kwargs):
            if not first.is_set():
                first.set()
                self.lock_database()
            return original(*args, **kwargs)
        self.stack.enter_context(patch.object(audit, name, side_effect=collide))
        campaign = await self.start()
        await self.until(lambda: f"SQLITE_RETRY phase={phase}" in self.output.getvalue())
        self.assertFalse(campaign.done())
        self.assertEqual(audit.runtime_snapshot(self.ns)["live_workers"], 4)
        self.assertEqual(self.row(reg)["state"], "running")
        self.blockers[0].rollback()
        await self.until(lambda: self.row(reg)["state"] == "completed")
        self.assertEqual(self.findings(reg), 1)
        self.assertEqual(audit.progress(self.ns)["new_evidence_based_findings"], 1)
        self.assertEqual(analyzer.call_count, 1)
        self.assertEqual(checker.call_count, 1)

    async def test_actual_inspection_persist_lock_retries_without_terminal_failure(self):
        await self.phase_lock("persist")

    async def test_actual_result_mark_lock_retries_without_double_count(self):
        await self.phase_lock("save")

    async def test_supervisor_restarts_one_service_and_heartbeat_survives_read_error(self):
        reg = self.purchase()
        original_claim, original_progress = audit.claim_next, audit.progress
        claimed = threading.Event()
        progress_failed = threading.Event()
        def claim(ns):
            if not claimed.is_set():
                claimed.set()
                raise RuntimeError("one unexpected worker failure")
            return original_claim(ns)
        def progress(ns):
            if ns.get("_v332_runtime", {}).get("status") == "running" and not progress_failed.is_set():
                progress_failed.set()
                raise sqlite3.OperationalError("database is locked")
            return original_progress(ns)
        self.stack.enter_context(patch.object(audit, "claim_next", side_effect=claim))
        self.stack.enter_context(patch.object(audit, "progress", side_effect=progress))
        campaign = await self.start()
        # The executor sets progress_failed before the heartbeat coroutine
        # observes its exception and writes the diagnostic log.
        await self.until(lambda: progress_failed.is_set() and self.row(reg)["state"] == "completed"
                         and audit.runtime_snapshot(self.ns)["live_workers"] == 4
                         and '"progress_error"' in self.output.getvalue())
        self.assertFalse(campaign.done())
        self.assertEqual(self.findings(reg), 1)
        self.assertIn("SERVICE_RESTART service=worker-", self.output.getvalue())
        self.assertIn('"progress_error"', self.output.getvalue())

    async def test_cancelled_late_claim_returns_its_reservation(self):
        reg = self.purchase()
        audit.initialize(self.ns)
        entered, release = threading.Event(), threading.Event()
        original = audit.claim_next
        def delayed_claim(ns):
            entered.set()
            if not release.wait(3): raise AssertionError("test did not release claim")
            return original(ns)
        with patch.object(audit, "claim_next", side_effect=delayed_claim):
            task = asyncio.create_task(audit.claim_safely(self.ns))
            try:
                await self.until(entered.is_set)
                task.cancel()
                await self.real_sleep(.02)
                self.assertFalse(task.done())
                release.set()
                with self.assertRaises(asyncio.CancelledError):
                    await asyncio.wait_for(task, 3)
            finally:
                release.set()
                await asyncio.gather(task, return_exceptions=True)
        self.assertEqual(self.row(reg)["state"], "pending")
        self.assertEqual(original(self.ns), reg)

    async def test_cancel_during_persist_waits_for_commit_before_releasing_reg(self):
        reg = self.purchase()
        entered, release = threading.Event(), threading.Event()
        original = audit.persist
        def delayed_persist(*args, **kwargs):
            entered.set()
            if not release.wait(3): raise AssertionError("test did not release persist")
            return original(*args, **kwargs)
        self.stack.enter_context(patch.object(audit, "persist", side_effect=delayed_persist))
        campaign = await self.start()
        try:
            await self.until(entered.is_set)
            campaign.cancel()
            await self.real_sleep(.025)
            self.assertFalse(campaign.done())
            self.assertEqual(self.row(reg)["state"], "running")
        finally:
            release.set()
        await asyncio.wait_for(asyncio.gather(campaign, return_exceptions=True), 3)
        self.assertEqual(self.row(reg)["state"], "pending")
        self.assertEqual(self.findings(reg), 1)
        self.assertEqual(audit.progress(self.ns)["new_evidence_based_findings"], 1)

    async def cancel_after_queue_commit(self, deferred):
        reg = self.purchase()
        if deferred:
            extra = self.document(reg, "Временно не прочитанное приложение.")
            original_extract = audit.document_text
            def extract(path, extractor):
                return "" if Path(path) == extra else original_extract(path, extractor)
            self.stack.enter_context(patch.object(audit, "document_text", side_effect=extract))
        name = "defer_retry" if deferred else "mark"
        original = getattr(audit, name)
        committed, release = threading.Event(), threading.Event()
        def completed_write(*args, **kwargs):
            result = original(*args, **kwargs)
            committed.set()
            if not release.wait(3): raise AssertionError("test did not release completed write")
            return result
        self.stack.enter_context(patch.object(audit, name, side_effect=completed_write))
        campaign = await self.start()
        try:
            await self.until(committed.is_set)
            before = self.row(reg)
            campaign.cancel()
            await self.real_sleep(.02)
            self.assertFalse(campaign.done())
        finally:
            release.set()
        await asyncio.wait_for(asyncio.gather(campaign, return_exceptions=True), 3)
        self.assertEqual(self.row(reg), before)
        self.assertEqual(before["state"], "pending" if deferred else "completed")
        self.assertEqual(self.findings(reg), 1)
        self.assertEqual(audit.progress(self.ns)["new_evidence_based_findings"], 1)

    async def test_cancellation_preserves_committed_completion(self):
        await self.cancel_after_queue_commit(False)

    async def test_cancellation_preserves_committed_partial_retry(self):
        await self.cancel_after_queue_commit(True)

    async def test_partial_result_keeps_finding_and_counter_through_delayed_retry(self):
        reg = self.purchase()
        extra = self.document(reg, "Дополнительные условия поставки.")
        readable = threading.Event()
        original = audit.document_text
        def extract(path, extractor):
            if Path(path) == extra and not readable.is_set(): return ""
            return original(path, extractor)
        self.stack.enter_context(patch.object(audit, "document_text", side_effect=extract))
        await self.start()
        await self.until(lambda: self.row(reg) is not None and self.row(reg)["retry_count"] == 1)
        first = self.row(reg)
        self.assertEqual(first["state"], "pending")
        self.assertGreater(first["next_retry_at"], time.time())
        self.assertEqual(first["documents_read"], 1)
        self.assertEqual(self.findings(reg), 1)
        self.assertEqual(audit.progress(self.ns)["new_evidence_based_findings"], 1)
        self.assertIsNone(audit.claim_next(self.ns))
        self.assertEqual(audit.refresh_changed(self.ns), 0)
        readable.set()  # Same file/status/id; only a transient extractor recovers.
        await self.until(lambda: self.row(reg)["state"] == "completed")
        self.assertEqual(self.row(reg)["documents_read"], 2)
        self.assertEqual(self.row(reg)["retry_count"], 0)
        self.assertEqual(self.findings(reg), 1)
        self.assertEqual(audit.progress(self.ns)["new_evidence_based_findings"], 1)
        self.assertEqual(audit.refresh_changed(self.ns), 0)
        with self.connect() as con:
            con.execute("UPDATE v332_audit_queue SET state='pending' WHERE reg=?", (reg,))
        notify_evidence(self.ns)
        await self.until(lambda: self.row(reg)["state"] == "completed")
        self.assertEqual(self.findings(reg), 1)
        self.assertEqual(audit.progress(self.ns)["new_evidence_based_findings"], 1)

    async def test_downloaded_no_text_original_is_retried_without_row_or_file_change(self):
        reg = self.purchase(status="downloaded_no_text")
        readable = threading.Event()
        original = audit.document_text
        def extract(path, extractor):
            return original(path, extractor) if readable.is_set() else ""
        self.stack.enter_context(patch.object(audit, "document_text", side_effect=extract))
        await self.start()
        await self.until(lambda: self.row(reg) is not None and self.row(reg)["retry_count"] == 1)
        self.assertEqual(self.row(reg)["state"], "pending")
        self.assertNotEqual(self.row(reg)["audit_revision"], audit.RULE_REVISION)
        self.assertEqual(self.findings(reg), 0)
        snapshot = audit.evidence_version(self.ns, reg)
        readable.set()
        await self.until(lambda: self.row(reg)["state"] == "completed")
        self.assertEqual(audit.evidence_version(self.ns, reg), snapshot)
        self.assertEqual(self.findings(reg), 1)
        with self.connect() as con:
            self.assertEqual(con.execute("SELECT status FROM docs WHERE reg_number=?", (reg,)).fetchone()[0], "downloaded_no_text")

    async def test_timeout_is_a_delayed_retry_with_explicit_reason(self):
        reg = self.purchase()
        original = audit.inspect_one
        attempts = 0
        async def timed_out_once(ns, number):
            nonlocal attempts
            attempts += 1
            if attempts == 1: raise asyncio.TimeoutError()
            return await original(ns, number)
        self.stack.enter_context(patch.object(audit, "inspect_one", side_effect=timed_out_once))
        await self.start()
        await self.until(lambda: self.row(reg) is not None and self.row(reg)["retry_count"] == 1)
        row = self.row(reg)
        self.assertEqual(row["state"], "pending")
        self.assertIn("600 секунд", row["error"])
        self.assertNotEqual(row["audit_revision"], audit.RULE_REVISION)
        self.assertIsNone(audit.claim_next(self.ns))
        await self.until(lambda: self.row(reg)["state"] == "completed")
        self.assertEqual(attempts, 2)
        self.assertEqual(self.findings(reg), 1)


if __name__ == "__main__":
    unittest.main()
