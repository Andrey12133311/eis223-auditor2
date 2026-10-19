"""Synthetic original files exercise the real extraction and existing rules.

No live procurement evidence, external requests or new legal thresholds.
"""
import asyncio
from contextlib import ExitStack, redirect_stdout
import hashlib
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from docx import Document
import v298_eis_only as incoming
import v311_verified_customer as customer
import v314_eis_metadata as metadata
import v332_all_procurements_audit as general
import v341_precise_legal as legal
import v342_thematic_queue as thematic

REG = "32616480000"
URL = "https://zakupki.gov.ru/223/filestore/public/1.0/download/fz223/file.html?uid="
P = {"reg": REG, "reg_number": REG, "title": "Поставка тестовых материалов",
     "customer": "ООО Тест", "method": "Открытый конкурс",
     "published_at": "2026-10-01", "published": "2026-10-01",
     "submission_start": "2026-10-01", "submission_end": "2026-10-30"}


def legacy_docx(path):
    """Faithful paragraph-then-table branch of the existing generic extractor."""
    document = Document(path)
    return "\n".join([p.text for p in document.paragraphs] +
                     [" | ".join(c.text for c in row.cells)
                      for table in document.tables for row in table.rows])


def verified_doc(text, name="Разъяснения.docx", uid="A" * 32):
    return {"name": name, "text": text, "status": "checked", "url": URL + uid,
            "local": "/data/" + name,
            "meta": {"source": "official_eis_original", "source_url": URL + uid,
                     "sha256": hashlib.sha256(text.encode()).hexdigest()}}


class IncomingEvidenceCoverage(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        self.stack.enter_context(patch.dict(os.environ, {"EIS223_DATA_DIR": str(self.root)}))
        self.stack.enter_context(patch.object(incoming, "stored_purchase", return_value=P.copy()))
        self.stack.enter_context(patch.object(customer, "read_verified", return_value={}))
        self.stack.enter_context(patch.object(customer, "read_verified_price", return_value={}))
        self.entries = []
        self.saved = []
        self.ns = {"extract_text": legacy_docx,
                   "analyze": lambda p, docs: legal.evaluate(p, docs)[0],
                   "build_checks": lambda p, detail, card, docs: legal.evaluate(p, docs)[1],
                   "save": lambda *args: self.saved.append(args)}

    def docx(self, missing_auxiliary=False):
        document = Document()
        # Moving both tables to the end makes the first request lose its
        # answer and the second own two answers. Ordered text has one late
        # reply and one timely reply, using the unchanged date validator.
        for request, answer in (("01.10.2026", "07.10.2026"),
                                ("08.10.2026", "08.10.2026")):
            document.add_paragraph("Дата поступления запроса разъяснения: " + request)
            cells = document.add_table(rows=1, cols=2).rows[0].cells
            cells[0].text = "Дата публикации ответа в ЕИС"
            cells[1].text = answer
        raw = io.BytesIO()
        document.save(raw)
        if not missing_auxiliary:
            return raw.getvalue()
        output = io.BytesIO()
        with zipfile.ZipFile(raw) as original, zipfile.ZipFile(output, "w") as package:
            for part in original.infolist():
                if part.filename != "word/styles.xml":
                    package.writestr(part, original.read(part))
        return output.getvalue()

    def store(self, data, name="Разъяснения.docx"):
        uid = format(len(self.entries) + 1, "032X")
        digest = hashlib.sha256(data).hexdigest()
        folder = self.root / "eis223_official_originals" / REG / uid
        folder.mkdir(parents=True)
        suffix = Path(name).suffix
        path = folder / (digest + suffix)
        path.write_bytes(data)
        (folder / "receipt.json").write_text(json.dumps({
            "reg_number": REG, "uid": uid, "extension": suffix,
            "sha256": digest, "bytes": len(data),
        }))
        self.entries.append({"url": URL + uid, "name": name})
        manifest = self.root / "eis223_public_browser_intake" / REG / "manifest.json"
        manifest.parent.mkdir(parents=True, exist_ok=True)
        manifest.write_text(json.dumps({"reg_number": REG,
            "source": "public_eis_223_browser", "documents": self.entries}))
        return path

    def audit(self, collect_only=False):
        return asyncio.run(incoming.offline_audit(self.ns, REG, collect_only=collect_only))

    def test_ordered_incoming_text_restores_existing_proven_rule(self):
        path = self.store(self.docx())
        original = path.read_bytes()
        legacy_findings, _ = legal.evaluate(P, [verified_doc(legacy_docx(path))])
        self.assertEqual(legacy_findings, [])  # Reproduce the old missed result.
        reply = self.audit()
        self.assertEqual(reply["findings"], 1)
        _, documents, findings, _, _ = self.saved[0]
        self.assertEqual(documents[0]["text"], metadata.document_text(path, legacy_docx))
        self.assertEqual(findings[0]["code"], "223-3.2-3-CLARIFICATION-RESPONSE-LATE")
        self.assertIn("2026-10-07", findings[0]["evidence"])
        self.assertEqual(path.read_bytes(), original)

    def test_archive_members_use_same_ordered_extractor(self):
        source = io.BytesIO()
        with zipfile.ZipFile(source, "w") as archive:
            archive.writestr("Разъяснения.docx", self.docx())
        path = self.store(source.getvalue(), "Документы.zip")
        reply = self.audit()
        self.assertEqual((reply["findings"], reply["checked_documents"]), (1, 1))
        _, documents, findings, _, _ = self.saved[0]
        self.assertEqual([d["status"] for d in documents], ["archive_processed", "checked"])
        self.assertEqual(documents[1]["meta"]["source"], "official_eis_archive_member")
        self.assertEqual(findings[0]["source_url"], self.entries[0]["url"])
        self.assertEqual(path.read_bytes(), source.getvalue())

    def test_missing_styles_does_not_drop_readable_original(self):
        path = self.store(self.docx(missing_auxiliary=True))
        with self.assertRaises(KeyError):
            legacy_docx(path)
        reply = self.audit()
        self.assertEqual((reply["checked_documents"], reply["findings"]), (1, 1))

    def test_one_unreadable_file_does_not_abort_other_originals(self):
        self.store(b"invalid docx package", "Повреждённый.docx")
        self.store(self.docx())
        reply = self.audit()
        self.assertEqual((reply["checked_documents"], reply["findings"]), (1, 1))
        documents = self.saved[0][1]
        self.assertEqual(documents[0]["status"], "downloaded_no_text")
        self.assertEqual(documents[0]["text"], "")
        self.assertTrue(documents[0]["error"])
        self.assertEqual(documents[1]["status"], "checked")

    def test_changed_bytes_do_not_reach_legal_engine(self):
        path = self.store(self.docx())
        raw = path.read_bytes()
        path.write_bytes(raw[:-1] + bytes([raw[-1] ^ 1]))
        with self.assertRaisesRegex(ValueError, "No readable original"):
            self.audit()
        self.assertEqual(self.saved, [])

    def test_collect_only_has_same_evidence_without_saving(self):
        path = self.store(self.docx())
        reply = self.audit(collect_only=True)
        self.assertEqual(reply["docs"][0]["text"], metadata.document_text(path, legacy_docx))
        self.assertEqual(self.saved, [])

    def test_log_counts_confirmed_checklist_episodes_without_duplicates(self):
        self.store(self.docx())
        def check_results(p, detail, card, docs):
            finding = legal.evaluate(p, docs)[0][0]
            check = {**finding, "item": finding["title"], "field_path": finding["code"],
                     "status": "violation"}
            return [check, check, {**check, "status": "check"}]
        self.ns["analyze"] = lambda p, docs: []
        self.ns["build_checks"] = check_results
        output = io.StringIO()
        with redirect_stdout(output):
            reply = self.audit()
        self.assertEqual(reply["findings"], 0)
        self.assertEqual(reply["checklist_violations"], 2)
        self.assertEqual(reply["confirmed_violations"], 1)
        self.assertIn("confirmed_violations=1", output.getvalue())


class ThematicChecklistCoverage(unittest.TestCase):
    def setUp(self):
        self.doc = verified_doc("Дата поступления запроса разъяснения: 01.10.2026. "
            "Дата публикации ответа в ЕИС: 07.10.2026.")
        self.finding = legal.evaluate(P, [self.doc])[0][0]
        self.check = {**self.finding, "status": "violation",
                      "item": self.finding["title"], "field_path": self.finding["code"]}
        for key in ("kind", "code", "title"):
            self.check.pop(key)
        self.card = {"status": "partial", "pages": {}, "summary": {"method": P["method"]}}
        self.purchase = {**P, "card_fields_json": json.dumps(self.card)}
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        self.stack.enter_context(patch.object(general, "existing", return_value=(self.purchase, [])))
        self.stack.enter_context(patch.object(general, "stored_docs", return_value=([self.doc], 0)))

    def review(self, findings=(), checks=(), categories=("dates",)):
        self.analyzer = Mock(return_value=list(findings))
        self.checker = Mock(return_value=list(checks))
        return asyncio.run(thematic.independent_review(
            {"analyze": self.analyzer, "build_checks": self.checker}, REG, categories))

    def test_confirmed_checklist_finding_is_included_with_original_identity(self):
        state, results = self.review(checks=[self.check])
        self.assertEqual((state, len(results)), ("completed", 1))
        self.assertEqual(results[0]["code"], self.finding["code"])
        self.assertEqual(results[0]["source_url"], self.doc["url"])
        self.assertEqual(results[0]["evidence"], self.finding["evidence"])
        self.assertEqual(self.checker.call_args.args[2], self.card)
        self.assertIs(self.analyzer.call_args.args[0], self.checker.call_args.args[0])
        self.assertIs(self.analyzer.call_args.args[1], self.checker.call_args.args[3])

    def test_analyzer_and_checklist_copies_are_one_episode(self):
        _, results = self.review(findings=[self.finding], checks=[self.check, self.check])
        self.assertEqual(len(results), 1)

    def test_risk_unknown_source_missing_law_and_empty_evidence_are_excluded(self):
        unsafe = [{**self.check, "status": state} for state in ("check", "risk", "ok")]
        unsafe += [{**self.check, "document": "unknown.docx", "source_url": "https://example.org"},
                   {**self.check, "law": ""}, {**self.check, "evidence": ""}]
        self.assertEqual(self.review(checks=unsafe), ("completed", []))
        self.assertEqual(self.review(checks=[self.check], categories=["national"]), ("completed", []))

    def test_different_dated_events_survive_shared_category(self):
        other = {**self.check, "evidence": self.check["evidence"].replace(
            "2026-10-01", "2026-10-08").replace("2026-10-06", "2026-10-13").replace(
            "2026-10-07", "2026-10-14")}
        _, results = self.review(checks=[self.check, other])
        self.assertEqual(len(results), 2)

    def test_missing_checklist_does_not_report_completed_full_audit(self):
        with self.assertRaisesRegex(RuntimeError, "checklist"):
            asyncio.run(thematic.independent_review({"analyze": lambda p, docs: []}, REG, ["dates"]))


if __name__ == "__main__":
    unittest.main()
