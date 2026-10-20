"""OOXML whitespace must not destroy explicit event IDs or date labels."""
import hashlib
import os
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import v314_eis_metadata as metadata
import v347_text_cache as cache
import test_date_label_pipeline_parity as pipeline
import test_official_publication_evidence as saved


class DocxEventSeparators(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix='eis-docx-separators-')
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        environment = patch.dict(os.environ, {'EIS223_DATA_DIR': str(self.root)})
        environment.start()
        self.addCleanup(environment.stop)
        directory = patch.object(cache.tempfile, 'gettempdir', return_value=str(self.root))
        directory.start()
        self.addCleanup(directory.stop)

    def save(self, doc):
        path = self.root / 'original.docx'
        doc.save(path)
        return path

    def extract(self, path):
        return metadata.document_text(path, lambda unused: self.fail('DOCX fallback used'))

    def test_manual_breaks_carriage_returns_and_tabs_keep_field_boundaries(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run('Номер запроса: 1').add_break()
        paragraph.add_run('Дата поступления запроса:').add_tab()
        paragraph.add_run('06.10.2026')._r.append(OxmlElement('w:cr'))
        paragraph.add_run('Дата публикации ответа в ЕИС: 13.10.2026')
        self.assertEqual(self.extract(self.save(doc)),
            'Номер запроса: 1\nДата поступления запроса:\t06.10.2026\n'
            'Дата публикации ответа в ЕИС: 13.10.2026')

    def test_adjacent_runs_join_without_added_word_spaces(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run('разъяс')
        paragraph.add_run('нений')
        self.assertEqual(self.extract(self.save(doc)), 'разъяснений')

    def test_paragraph_and_table_order_remains_source_order(self):
        doc = Document()
        doc.add_paragraph('До таблицы')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Дата поступления\nзапроса'
        table.cell(0, 1).text = '06.10.2026'
        doc.add_paragraph('После таблицы')
        self.assertEqual(self.extract(self.save(doc)),
                         'До таблицы\nДата поступления запроса | 06.10.2026\nПосле таблицы')

    def test_nested_paragraph_text_is_visited_once_in_xml_order(self):
        doc = Document()
        paragraph = doc.add_paragraph('Перед')
        box = OxmlElement('w:txbxContent')
        inside = OxmlElement('w:p')
        run = OxmlElement('w:r')
        text = OxmlElement('w:t')
        text.text = 'Внутренний текст'
        run.append(text)
        inside.append(run)
        box.append(inside)
        paragraph._p.append(box)
        paragraph.add_run('После')
        result = self.extract(self.save(doc))
        self.assertEqual(result, 'Перед\nВнутренний текст\nПосле')
        self.assertEqual(result.count('Внутренний текст'), 1)

    def test_old_concatenated_cache_is_reextracted_under_new_namespace(self):
        doc = Document()
        doc.add_paragraph('Номер запроса: 1\nДата поступления запроса: 06.10.2026')
        path = self.save(doc)
        legacy = 'Номер запроса: 1Дата поступления запроса: 06.10.2026'
        self.assertEqual(cache.cached_extract(path, lambda unused: legacy,
            namespace='eis223-docx-evidence-order-v381', data_root=self.root), legacy)
        with patch.object(metadata, '_document_text_uncached',
                          wraps=metadata._document_text_uncached) as parser:
            self.assertEqual(self.extract(path), legacy.replace('1Дата', '1\nДата'))
            self.assertEqual(self.extract(path), legacy.replace('1Дата', '1\nДата'))
            self.assertEqual(parser.call_count, 1)
        self.assertEqual(cache.cached_extract(path, lambda unused: self.fail('Legacy cache lost'),
            namespace='eis223-docx-evidence-order-v381', data_root=self.root), legacy)


class DocxSeparatorReceiptPipeline(unittest.TestCase):
    setUp = saved.OfficialPublicationEvidenceTests.setUp
    connect = saved.OfficialPublicationEvidenceTests.connect
    make_xml = saved.OfficialPublicationEvidenceTests.make_xml
    write_archive = saved.OfficialPublicationEvidenceTests.write_archive
    write_original = saved.OfficialPublicationEvidenceTests.write_original
    collect = saved.OfficialPublicationEvidenceTests.collect
    seed_collected = pipeline.DateLabelPipelineParity.seed_collected
    assert_both_paths = pipeline.DateLabelPipelineParity.assert_both_paths

    def test_break_and_tab_numbered_clarification_joins_only_verified_originals(self):
        answer_uid = 'A' * 32
        # The generic fixture starts with a protocol attachment under this UID.
        # Give this test one consistent documentary role for each original.
        shutil.rmtree(self.root / 'eis223_official_archive_inbox' / saved.REG)
        self.write_archive(self.make_xml('purchaseNotice', version=2,
            filename='Запрос разъяснений № 1.docx'))
        self.write_original(
            'Номер запроса:\t1\n'
            'Дата поступления запроса о даче разъяснений положений документации:\t'
            '06.10.2026\nЗапрос участника закупки принят заказчиком.')
        self.write_archive(self.make_xml('purchaseNotice', version=3, uid=answer_uid,
            filename='Разъяснение на запрос № 1.docx'))
        self.write_original(
            'Номер запроса:\t1\n'
            'Дата размещения в ЕИС разъяснений положений документации:\t13.10.2026\n'
            'Предоставлен ответ на поступивший запрос участника закупки.', uid=answer_uid)
        collected = self.seed_collected()
        self.assertEqual(len(collected['docs']), 2)
        for doc in collected['docs']:
            self.assertIn('Номер запроса:\t1\nДата', doc['text'])
            self.assertEqual(doc['meta']['sha256'],
                             hashlib.sha256(Path(doc['local']).read_bytes()).hexdigest())
        found = self.assert_both_paths(1, seed=False)
        self.assertEqual(found[0]['code'], '223-3.2-3-CLARIFICATION-CROSS-DOC-LATE')
        self.assertIn('Номер запроса: 1.', found[0]['evidence'])


if __name__ == '__main__':
    unittest.main()
