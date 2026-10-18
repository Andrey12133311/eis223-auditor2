"""Local extraction reuse regressions; no network or live EIS requests."""
import concurrent.futures
import json
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch

import v347_text_cache as cache
import v314_eis_metadata as metadata


class ExtractionThroughput(unittest.TestCase):
    def setUp(self):
        self.home = tempfile.TemporaryDirectory()
        self.root = Path(self.home.name)
        self.root_patch = patch('v295_eis_originals.root', return_value=self.root)
        self.root_patch.start()
        # Isolate cache across test runs.
        self.tmp_patch = patch.object(cache.tempfile, 'gettempdir', return_value=self.home.name)
        self.tmp_patch.start()

    def tearDown(self):
        self.tmp_patch.stop()
        self.root_patch.stop()
        self.home.cleanup()

    def test_format_parser_and_bytes_are_separate(self):
        pdf, txt = self.root/'evidence.pdf', self.root/'evidence.txt'
        pdf.write_bytes(b'identical bytes')
        txt.write_bytes(pdf.read_bytes())
        def extract(p): return 'format ' + Path(p).suffix
        self.assertEqual(cache.cached_extract(pdf,extract),'format .pdf')
        self.assertEqual(cache.cached_extract(txt,extract),'format .txt')
        self.assertEqual(cache.cached_extract(pdf,lambda _: 'different parser',namespace='other'),'different parser')
        pdf.write_bytes(b'replaced bytes')
        self.assertEqual(cache.cached_extract(pdf,lambda _: 'new evidence'),'new evidence')

    def test_empty_errors_and_outside_root_never_cached(self):
        p=self.root/'sample.pdf'; p.write_bytes(b'valid source')
        self.assertEqual(cache.cached_extract(p, lambda _: ''),'')
        self.assertEqual(cache.cached_extract(p, lambda _: 'retry success'),'retry success')
        with tempfile.TemporaryDirectory(dir="/tmp") as other:
            outside=Path(other)/'sample.pdf';outside.write_bytes(b'outside')
            self.assertEqual(cache.cached_extract(outside,lambda _: 'first'),'first')
            self.assertEqual(cache.cached_extract(outside,lambda _: 'second'),'second')

    def test_singleflight_and_bounded_locks(self):
        p=self.root/'sample.pdf';p.write_bytes(b'one verified input')
        calls=[]
        def extract(_):
            calls.append(1);time.sleep(.02);return 'the unchanged evidence'
        with concurrent.futures.ThreadPoolExecutor(max_workers=8) as pool:
            values=list(pool.map(lambda _:cache.cached_extract(p,extract),range(16)))
        self.assertEqual(len(calls),1)
        self.assertEqual(len(set(values)),1)
        self.assertEqual(len(cache._locks),128)

    def test_replacement_during_extraction_does_not_poison(self):
        p=self.root/'sample.pdf';p.write_bytes(b'original input')
        def change(_):
            p.write_bytes(b'replacement input');return 'changed evidence'
        cache.cached_extract(p,change)
        p.write_bytes(b'original input')
        self.assertEqual(cache.cached_extract(p,lambda _: 'original evidence'),'original evidence')

    def test_docx_order_and_warm_reuse_benchmark(self):
        from docx import Document
        doc=Document()
        doc.add_paragraph('Начало документа')
        for i in range(100):
            doc.add_paragraph('Условия закупки '+str(i))
            cells=doc.add_table(rows=1,cols=2).rows[0].cells
            cells[0].text=str(i)+'.';cells[1].text='Заказчик\nООО Проверка'
        doc.add_paragraph('Окончание документа')
        p=self.root/'evidence.docx';doc.save(p)
        fallback=lambda _: 'must not replace ordered DOCX parser'
        t=time.perf_counter()
        original=[metadata._document_text_uncached(p,fallback) for _ in range(8)]
        baseline=time.perf_counter()-t
        t=time.perf_counter()
        result=[metadata.document_text(p,fallback) for _ in range(8)]
        cached=time.perf_counter()-t
        self.assertEqual(original,result)
        self.assertLess(result[0].index('Начало документа'),result[0].index('0. | Заказчик'))
        self.assertLess(result[0].index('99. | Заказчик'),result[0].index('Окончание документа'))
        print('LOCAL_DOCX_BENCHMARK '+json.dumps({'repeats':8,'uncached_seconds':round(baseline,4),
            'cached_seconds':round(cached,4),'speedup':round(baseline/cached,2),
            'text_identical':True}),flush=True)

if __name__ == '__main__':unittest.main()
