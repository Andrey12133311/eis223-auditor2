"""Real DOCX relationship failures must not discard an intact original body."""
import asyncio
import contextlib
import hashlib
import io
import json
from pathlib import Path
import sqlite3
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from docx import Document
import v314_eis_metadata as metadata
import v332_all_procurements_audit as audit


REG = "32616467075"
UID = "A" * 32
URL = "https://zakupki.gov.ru/223/filestore/public/1.0/download/fz223/file.html?uid=" + UID
CUSTOMER = "Государственное автономное учреждение «Агентство»"
TITLE = "Оказание услуг по организации мероприятий"


class DocxMetadataRecoveryTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.home = Path(self.temporary.name)
        self.env = patch.dict("os.environ", {"EIS223_DATA_DIR": str(self.home)})
        self.env.start()
        self.addCleanup(self.env.stop)
        self.database = self.home / "test.sqlite"
        with self.connect() as db:
            db.executescript("""
                CREATE TABLE purchases(reg_number TEXT PRIMARY KEY,customer TEXT);
                CREATE TABLE docs(id INTEGER PRIMARY KEY,reg_number TEXT,name TEXT,
                  url TEXT,doc_type TEXT,published_at TEXT,local_path TEXT,
                  metadata_json TEXT,status TEXT,final_url TEXT);
                CREATE TABLE findings(id INTEGER PRIMARY KEY,evidence TEXT);
                INSERT INTO findings VALUES(1,'existing finding');
            """)
            db.execute("INSERT INTO purchases VALUES(?,NULL)", (REG,))
        self.ns = {"conn": self.connect, "extract_text": lambda p: Path(p).read_text()}

    def connect(self):
        db = sqlite3.connect(self.database)
        db.row_factory = sqlite3.Row
        return db

    def docx(self, missing="word/styles.xml", replacement_body=None):
        source = self.home / "source.docx"
        document = Document()
        document.add_paragraph("Сведения о закупке")
        table = document.add_table(rows=3, cols=2)
        for row, pair in zip(table.rows, [
            ("Наименование заказчика", CUSTOMER),
            ("Наименование закупки", TITLE),
            ("ИНН заказчика", "6501309746"),
        ]):
            for cell, value in zip(row.cells, pair):
                cell.text = value
        document.add_paragraph("Последний абзац документа")
        document.save(source)
        output = io.BytesIO()
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(output, "w") as broken:
            for part in original.infolist():
                if part.filename == missing:
                    continue
                data = original.read(part)
                if part.filename == "word/document.xml" and replacement_body is not None:
                    data = replacement_body
                broken.writestr(part, data)
        return output.getvalue()

    def store(self, data, suffix=".docx", source_url=True):
        digest = hashlib.sha256(data).hexdigest()
        folder = self.home / "eis223_official_originals" / REG / UID
        folder.mkdir(parents=True)
        path = folder / (digest + suffix)
        path.write_bytes(data)
        (folder / "receipt.json").write_text(json.dumps({
            "reg_number": REG, "uid": UID, "extension": suffix,
            "sha256": digest, "bytes": len(data),
        }))
        manifest = self.home / "eis223_public_browser_intake" / REG / "manifest.json"
        manifest.parent.mkdir(parents=True)
        manifest.write_text(json.dumps({
            "reg_number": REG, "source": "public_eis_223_browser",
            "documents": [{"url": URL, "name": "Извещение" + suffix}],
        }))
        meta = {"source": "official_eis_original", "sha256": digest}
        if source_url:
            meta["source_url"] = URL
        with self.connect() as db:
            db.execute("INSERT INTO docs VALUES(1,?,?,?,?,?,?,?,?,?)", (
                REG, "Извещение" + suffix, URL, "notice", "2026-10-18",
                str(path), json.dumps(meta), "checked", URL,
            ))
            row = dict(db.execute("SELECT * FROM docs").fetchone())
        return path, row

    def audit_docs(self, row):
        return asyncio.run(audit.stored_docs(self.ns, REG, [row]))

    def test_dangling_styles_reproduces_python_docx_keyerror(self):
        path, _ = self.store(self.docx())
        with self.assertRaisesRegex(KeyError, "word/styles.xml"):
            Document(path)
        with self.assertLogs(metadata.LOG, level="WARNING") as logs:
            text = metadata._document_text_uncached(path, self.ns["extract_text"])
        self.assertEqual(text, metadata.document_text(path, self.ns["extract_text"]))
        self.assertIn("word/styles.xml", logs.output[0])
        self.assertLess(text.index("Сведения о закупке"), text.index(CUSTOMER))
        self.assertLess(text.index(CUSTOMER), text.index("Последний абзац документа"))
        self.assertEqual(metadata.parse_text(text)["title"], TITLE)

    def test_metadata_recovery_preserves_original_files_and_saved_rows(self):
        path, row = self.store(self.docx(), source_url=False)
        original = path.read_bytes()
        fields = metadata.stored_metadata(self.ns, REG)
        self.assertEqual(fields["customer"], CUSTOMER)
        self.assertEqual(fields["customer_inn"], "6501309746")
        self.assertEqual(fields["title"], TITLE)
        self.assertEqual(fields["source_url"], URL)
        self.assertEqual(path.read_bytes(), original)
        with self.connect() as db:
            self.assertEqual(dict(db.execute("SELECT * FROM docs").fetchone()), row)
            self.assertEqual(db.execute("SELECT evidence FROM findings").fetchone()[0],
                             "existing finding")

    def test_audit_recovers_and_verifies_legacy_row_url_using_real_receipt(self):
        _, row = self.store(self.docx(), source_url=False)
        documents, unreadable = self.audit_docs(row)
        self.assertEqual(unreadable, 0)
        self.assertEqual(len(documents), 1)
        self.assertTrue(audit.is_official(documents[0]))
        self.assertIn(TITLE, documents[0]["text"])

    def test_missing_main_document_is_unreadable(self):
        _, row = self.store(self.docx(missing="word/document.xml"))
        self.assertEqual(metadata.stored_metadata(self.ns, REG), {})
        documents, unreadable = self.audit_docs(row)
        self.assertEqual((documents, unreadable), ([], 1))

    def test_corrupt_main_document_is_unreadable(self):
        _, row = self.store(self.docx(replacement_body=b"<invalid"))
        self.assertEqual(metadata.stored_metadata(self.ns, REG), {})
        documents, unreadable = self.audit_docs(row)
        self.assertEqual((documents, unreadable), ([], 1))

    def test_corrupt_main_document_with_all_auxiliary_parts_is_unreadable(self):
        _, row = self.store(self.docx(missing=None, replacement_body=b"<invalid"))
        self.assertEqual(metadata.stored_metadata(self.ns, REG), {})
        self.assertEqual(self.audit_docs(row), ([], 1))

    def test_valid_docx_keeps_existing_extraction_and_provenance(self):
        _, row = self.store(self.docx(missing=None))
        fields = metadata.stored_metadata(self.ns, REG)
        self.assertEqual((fields["customer"], fields["title"]), (CUSTOMER, TITLE))
        documents, unreadable = self.audit_docs(row)
        self.assertEqual(unreadable, 0)
        self.assertTrue(audit.is_official(documents[0]))

    def test_main_document_with_entity_declaration_is_not_recovered(self):
        xml = b'''<!DOCTYPE w:document [<!ENTITY injected "invented">]>
          <w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
          <w:body><w:p><w:r><w:t>&injected;</w:t></w:r></w:p></w:body></w:document>'''
        _, row = self.store(self.docx(replacement_body=xml))
        self.assertEqual(metadata.stored_metadata(self.ns, REG), {})
        self.assertEqual(self.audit_docs(row), ([], 1))

    def test_missing_or_modified_original_does_not_become_verified(self):
        path, row = self.store(self.docx())
        data = path.read_bytes()
        path.write_bytes(b"X" + data[1:])
        self.assertEqual(metadata.stored_metadata(self.ns, REG), {})
        self.assertEqual(self.audit_docs(row), ([], 1))
        path.unlink()
        self.assertEqual(metadata.stored_metadata(self.ns, REG), {})
        self.assertEqual(self.audit_docs(row), ([], 1))

    def test_missing_receipt_cannot_authenticate_readable_body(self):
        path, row = self.store(self.docx())
        (path.parent / "receipt.json").unlink()
        self.assertEqual(metadata.stored_metadata(self.ns, REG), {})
        documents, unreadable = self.audit_docs(row)
        self.assertEqual(unreadable, 0)
        self.assertEqual(documents[0]["meta"]["source"], "existing_local_unverified")
        self.assertFalse(audit.is_official(documents[0]))

    def test_extractor_keyerror_keeps_diagnostic_reason(self):
        _, row = self.store(b"original document bytes", suffix=".doc")
        def fail(_):
            raise KeyError("missing-parser-key")
        self.ns["extract_text"] = fail
        with self.assertLogs(metadata.LOG, level="WARNING") as logs:
            self.assertEqual(metadata.stored_metadata(self.ns, REG), {})
        self.assertIn("missing-parser-key", logs.output[0])
        self.assertIn("stage=text", logs.output[0])
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            self.assertEqual(self.audit_docs(row), ([], 1))
        self.assertIn("type=KeyError", output.getvalue())
        self.assertIn("stage=text", output.getvalue())
        self.assertIn("missing-parser-key", output.getvalue())


if __name__ == "__main__":
    unittest.main()
