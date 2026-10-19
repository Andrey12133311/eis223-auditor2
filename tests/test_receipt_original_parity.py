"""Real receipt files must reach both engines before the intake saves docs.

The originals and SQLite database are temporary. Existing clarification rules
establish the test findings; no remote I/O or new legal threshold is involved.
"""
import asyncio
import hashlib
import io
import json
from pathlib import Path
import sys
import time
import unittest
from unittest.mock import patch
import zipfile

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import audit_source_version as sources
import v298_eis_only as incoming
import v332_all_procurements_audit as general
import v341_precise_legal as legal
import v342_thematic_queue as thematic
import v345_thematic_accelerator as fast
from test_general_audit_recovery import DatabaseFixture, LATE_REPLY, URL


LATE_REPLY_B = ("Дата поступления запроса разъяснения: 08.10.2026.\n"
                "Дата публикации ответа в ЕИС: 14.10.2026.")
TIMELY_REPLY = ("Дата поступления запроса разъяснения: 01.10.2026.\n"
                "Дата публикации ответа в ЕИС: 02.10.2026.")


def docx_bytes(text):
    document = Document()
    document.add_paragraph(text)
    encoded = io.BytesIO()
    document.save(encoded)
    return encoded.getvalue()


class ReceiptOriginalParity(DatabaseFixture, unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        super().setUp()
        state, tick = dict(thematic.WORKER_STATE), thematic.tick

        def restore():
            thematic.WORKER_STATE.clear()
            thematic.WORKER_STATE.update(state)
            thematic.tick = tick

        self.addCleanup(restore)
        thematic.WORKER_STATE.clear()
        self.stack.enter_context(patch.object(thematic, "storage_health", return_value={"blocked": False}))
        self.stack.enter_context(patch.object(fast, "_request_missing_originals", return_value=0))
        thematic.init(self.ns)
        fast.install(self.ns)
        self.batch = "b" * 24
        with self.connect() as con:
            con.execute("INSERT INTO thematic_batches(id,categories,created_at,status,legal_forms,date_from,date_to) "
                        "VALUES(?,? ,?,'active',?, '', '')",
                        (self.batch, '["dates"]', thematic.now(), '["ООО"]'))

    def receipt_purchase(self, number=1):
        reg = self.purchase(number)
        # Reproduce the interval after the receipt commit and before docs save.
        with self.connect() as con:
            con.execute("DELETE FROM docs WHERE reg_number=?", (reg,))
        return reg

    def put_original(self, reg, data, name, uid=None):
        entries = self.originals.setdefault(reg, [])
        uid = uid or f"{len(entries) + 1:032X}"
        extension = Path(name).suffix
        digest = hashlib.sha256(data).hexdigest()
        folder = self.folder / "eis223_official_originals" / reg / uid
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / (digest + extension)
        path.write_bytes(data)
        (folder / "receipt.json").write_text(json.dumps({
            "reg_number": reg, "uid": uid, "extension": extension,
            "sha256": digest, "bytes": len(data)}))
        entry = {"url": URL + uid, "name": name}
        entries[:] = [item for item in entries if item["url"] != entry["url"]] + [entry]
        manifest = self.folder / "eis223_public_browser_intake" / reg / "manifest.json"
        manifest.parent.mkdir(parents=True, exist_ok=True)
        manifest.write_text(json.dumps({"reg_number": reg,
            "source": "public_eis_223_browser", "documents": entries}))
        return path

    def themed(self, reg):
        with self.connect() as con:
            item = dict(con.execute("SELECT * FROM thematic_items WHERE batch_id=? AND reg=?",
                                    (self.batch, reg)).fetchone())
            findings = [dict(row) for row in con.execute(
                "SELECT * FROM thematic_findings WHERE batch_id=? AND reg=?", (self.batch, reg))]
        return item, findings

    async def finish_general(self, reg):
        state, detail = await general.inspect_one(self.ns, reg)
        self.assertEqual(state, "completed")
        self.assertFalse(detail["source_incomplete"])
        general.mark(self.ns, reg, state,
                     {key: value for key, value in detail.items() if key != "source_incomplete"})
        return detail

    async def test_receipt_only_original_has_same_confirmed_finding_in_both_audits(self):
        reg = self.receipt_purchase()
        before = {path: path.read_bytes() for path in self.contents}
        general.initialize(self.ns)
        self.assertEqual(general.claim_next(self.ns), reg)
        first = await self.finish_general(reg)
        self.assertEqual((first["verified_original_documents"], first["new_findings"]), (1, 1))
        self.assertEqual(await fast.faster_tick(self.ns), 1)
        item, findings = self.themed(reg)
        self.assertEqual((item["status"], len(findings)), ("completed", 1))
        self.assertEqual(findings[0]["code"], "223-3.2-3-CLARIFICATION-RESPONSE-LATE")
        with self.connect() as con:
            general_finding = dict(con.execute("SELECT * FROM findings WHERE reg_number=?", (reg,)).fetchone())
            self.assertEqual(con.execute("SELECT COUNT(*) FROM docs").fetchone()[0], 0)
        for field in ("code", "evidence", "document", "source_url"):
            self.assertEqual(findings[0][field], general_finding[field])
        self.assertEqual(fast._prepare(self.ns), [])
        self.assertEqual(general.refresh_changed(self.ns), 0)
        self.assertEqual((await self.finish_general(reg))["new_findings"], 0)
        self.assertEqual(general.progress(self.ns)["new_evidence_based_findings"], 1)
        self.assertEqual({path: path.read_bytes() for path in before}, before)

    async def test_real_general_worker_completes_receipt_only_original_without_docs_save(self):
        reg = self.receipt_purchase()
        real_sleep = asyncio.sleep
        async def fast_start(seconds, *args, **kwargs):
            return await real_sleep(0 if seconds == 14 else seconds, *args, **kwargs)
        self.stack.enter_context(patch.object(general.asyncio, "sleep", side_effect=fast_start))
        self.stack.enter_context(patch.object(general, "QUEUE_POLL_SECONDS", .02))
        general.install(self.ns)
        await self.ns["app"].startups[0]()
        worker = self.ns["_v332_campaign_worker"]
        try:
            deadline = time.monotonic() + 5
            while time.monotonic() < deadline:
                if self.row(reg) and self.row(reg)["state"] == "completed":
                    break
                await real_sleep(.01)
            self.assertEqual(self.row(reg)["state"], "completed", self.output.getvalue())
            self.assertEqual((self.row(reg)["documents_read"], self.row(reg)["new_findings"]), (1, 1))
            self.assertEqual(general.runtime_snapshot(self.ns)["live_workers"], 4)
            with self.connect() as con:
                self.assertEqual(con.execute("SELECT COUNT(*) FROM docs").fetchone()[0], 0)
        finally:
            worker.cancel()
            await asyncio.wait_for(asyncio.gather(worker, return_exceptions=True), 5)

    def test_receipt_claim_reaches_later_page_without_holding_a_writer_over_files(self):
        reg = self.receipt_purchase(100)
        with self.connect() as con:
            for number in range(1, general.RECEIPT_CLAIM_PAGE + 1):
                missing = str(32619000000 + number)
                self.regs.append(missing)
                con.execute("INSERT INTO purchases(reg_number,customer,last_seen) VALUES(?,'ООО Тест','2026-10-19')",
                            (missing,))
        general.initialize(self.ns)
        self.assertIsNone(general.claim_next(self.ns))
        self.assertEqual(general.claim_next(self.ns), reg)
        general.return_running(self.ns, reg)
        original_ready = general.originals_ready
        def taken_during_read(number):
            ready = original_ready(number)
            if ready:
                with self.connect() as writer:
                    writer.execute("UPDATE v332_audit_queue SET state='running',updated_at='other-worker' WHERE reg=?", (number,))
                    writer.commit()
            return ready
        with patch.object(general, "originals_ready", side_effect=taken_during_read):
            self.assertIsNone(general.claim_next(self.ns))
            self.assertIsNone(general.claim_next(self.ns))
        self.assertEqual((self.row(reg)["state"], self.row(reg)["updated_at"]), ("running", "other-worker"))

    async def test_one_stored_archive_member_does_not_hide_the_other_member(self):
        reg = self.receipt_purchase()
        self.originals[reg] = []
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w") as container:
            container.writestr("Своевременный ответ.docx", docx_bytes(TIMELY_REPLY))
            container.writestr("Просроченный ответ.docx", docx_bytes(LATE_REPLY_B))
        path = self.put_original(reg, archive.getvalue(), "Разъяснения.zip")
        collected = await incoming.offline_audit(self.ns, reg, collect_only=True)
        member = next(doc for doc in collected["docs"] if doc["status"] == "checked")
        with self.connect() as con:
            con.execute("INSERT INTO docs(reg_number,name,url,doc_type,published_at,local_path,status,metadata_json,final_url) "
                        "VALUES(?,?,?,?,?,?,?,?,?)", (reg, member["name"], member["url"], "clarification", "",
                        member["local"], member["status"], json.dumps(member["meta"]), member["final_url"]))
        purchase, rows = general.existing(self.ns, reg)
        old_subset, unreadable = await general.stored_docs(self.ns, reg, rows)
        self.assertEqual((len(old_subset), unreadable), (1, 0))
        self.assertEqual(legal.evaluate(general.as_procurement(purchase), old_subset)[0], [])
        general.initialize(self.ns)
        detail = await self.finish_general(reg)
        self.assertEqual((detail["documents_read"], detail["new_findings"]), (2, 1))
        await fast.faster_tick(self.ns)
        item, findings = self.themed(reg)
        self.assertEqual((item["status"], len(findings)), ("completed", 1))
        self.assertIn("Просроченный ответ.docx", findings[0]["document"])
        self.assertEqual(findings[0]["source_url"], self.originals[reg][0]["url"])
        self.assertEqual(path.read_bytes(), archive.getvalue())
        with self.connect() as con:
            self.assertEqual(con.execute("SELECT COUNT(*) FROM docs").fetchone()[0], 1)

    async def test_archive_member_cache_is_restored_from_original_after_tamper_or_symlink(self):
        reg = self.receipt_purchase()
        self.originals[reg] = []
        archive = io.BytesIO()
        expected = docx_bytes(LATE_REPLY)
        with zipfile.ZipFile(archive, "w") as container:
            container.writestr("Разъяснение.docx", expected)
            container.writestr("../../escape.docx", docx_bytes(TIMELY_REPLY))
            symbolic = zipfile.ZipInfo("link.docx")
            symbolic.create_system = 3
            symbolic.external_attr = 0o120777 << 16
            container.writestr(symbolic, b"/tmp/should-not-read")
        original = self.put_original(reg, archive.getvalue(), "Разъяснения.zip")
        first = await incoming.offline_audit(self.ns, reg, collect_only=True)
        member = next(doc for doc in first["docs"] if doc["status"] == "checked")
        cached = Path(member["local"])
        other = self.folder / "unrelated.docx"
        other.write_bytes(docx_bytes(TIMELY_REPLY))
        for damage in ("bytes", "symlink"):
            with self.subTest(damage=damage):
                if damage == "bytes":
                    cached.write_bytes(other.read_bytes())
                else:
                    cached.unlink()
                    cached.symlink_to(other)
                repaired = await incoming.offline_audit(self.ns, reg, collect_only=True)
                checked = [doc for doc in repaired["docs"] if doc["status"] == "checked"]
                self.assertEqual(len(checked), 1)
                self.assertEqual(cached.read_bytes(), expected)
                self.assertFalse(cached.is_symlink())
                self.assertEqual(checked[0]["meta"]["sha256"], hashlib.sha256(cached.read_bytes()).hexdigest())
                self.assertEqual(len(legal.evaluate(repaired["p"], checked)[0]), 1)
                self.assertEqual(Document(other).paragraphs[0].text, TIMELY_REPLY)
                self.assertEqual(original.read_bytes(), archive.getvalue())
                self.assertFalse((self.folder / "escape.docx").exists())

    async def test_new_receipt_reopens_completed_audits_without_new_doc_id(self):
        reg = self.receipt_purchase()
        general.initialize(self.ns)
        await self.finish_general(reg)
        await fast.faster_tick(self.ns)
        previous = self.themed(reg)[0]
        self.assertEqual(fast._prepare(self.ns), [])
        self.put_original(reg, docx_bytes(LATE_REPLY_B), "Второе разъяснение.docx")
        prepared = fast._prepare(self.ns)
        self.assertEqual([item["reg"] for item in prepared], [reg])
        self.assertEqual(prepared[0]["ready"], 0)
        self.assertTrue(prepared[0]["originals_ready"])
        self.assertNotEqual(prepared[0]["source_version"], previous["last_source_version"])
        self.assertEqual(general.refresh_changed(self.ns), 1)
        self.assertEqual(general.claim_next(self.ns), reg)
        self.assertEqual((await self.finish_general(reg))["new_findings"], 1)
        await fast.faster_tick(self.ns)
        item, findings = self.themed(reg)
        self.assertEqual((item["status"], len(findings), item["last_doc_id"]), ("completed", 2, 0))
        self.assertEqual((self.findings(reg), general.progress(self.ns)["new_evidence_based_findings"]), (2, 2))
        self.assertEqual(fast._prepare(self.ns), [])
        self.assertEqual(general.refresh_changed(self.ns), 0)

    async def test_unreadable_receipt_original_preserves_finding_and_backs_off_until_repaired(self):
        reg = self.receipt_purchase()
        path = self.put_original(reg, b"damaged original docx", "Второе разъяснение.docx")
        uid = path.parent.name
        await fast.faster_tick(self.ns)
        item, findings = self.themed(reg)
        self.assertEqual((item["status"], item["retry_attempts"], len(findings)), ("retry", 1, 1))
        self.assertEqual(fast._prepare(self.ns), [])
        self.put_original(reg, docx_bytes(LATE_REPLY_B), "Второе разъяснение.docx", uid)
        await fast.faster_tick(self.ns)
        item, findings = self.themed(reg)
        self.assertEqual((item["status"], item["retry_attempts"], len(findings)), ("completed", 0, 2))
        self.assertEqual(fast._prepare(self.ns), [])

    async def test_archive_partial_extraction_is_retried_with_valid_findings_retained(self):
        reg = self.receipt_purchase()
        self.originals[reg] = []
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w") as container:
            container.writestr("Разъяснение.docx", docx_bytes(LATE_REPLY))
            container.writestr("Повреждённый документ.docx", b"invalid docx")
        self.put_original(reg, archive.getvalue(), "Разъяснения.zip")
        await fast.faster_tick(self.ns)
        item, findings = self.themed(reg)
        self.assertEqual((item["status"], item["retry_attempts"], len(findings)), ("retry", 1, 1))
        self.assertEqual(fast._prepare(self.ns), [])
        state, detail = await general.inspect_one(self.ns, reg)
        self.assertEqual((state, detail["new_findings"]), ("completed", 1))
        self.assertTrue(detail["source_incomplete"])

    async def test_failed_receipt_collection_does_not_finish_thematic_audit(self):
        reg = self.receipt_purchase()
        with patch.object(general, "offline_audit", side_effect=asyncio.TimeoutError()):
            await fast.faster_tick(self.ns)
        item, findings = self.themed(reg)
        self.assertEqual((item["status"], item["retry_attempts"], findings), ("retry", 1, []))
        self.assertEqual(fast._prepare(self.ns), [])

    async def test_tampered_or_unlisted_receipt_never_becomes_proven_violation(self):
        altered, unlisted = [self.receipt_purchase(number) for number in (1, 2)]
        for path in self.contents:
            if path.parent.parent.name == altered:
                raw = path.read_bytes()
                path.write_bytes(raw[:-1] + bytes([raw[-1] ^ 1]))
        manifest = self.folder / "eis223_public_browser_intake" / unlisted / "manifest.json"
        manifest.write_text(json.dumps({"reg_number": unlisted,
            "source": "public_eis_223_browser", "documents": []}))
        await fast.faster_tick(self.ns)
        self.assertEqual((self.themed(altered)[0]["status"], self.themed(altered)[1]), ("retry", []))
        self.assertEqual((self.themed(unlisted)[0]["status"], self.themed(unlisted)[1]), ("waiting_originals", []))

    async def test_tampered_supplemental_original_does_not_make_readable_subset_complete(self):
        reg = self.receipt_purchase()
        path = self.put_original(reg, docx_bytes(LATE_REPLY_B), "Второе разъяснение.docx")
        raw = path.read_bytes()
        path.write_bytes(raw[:-1] + bytes([raw[-1] ^ 1]))
        await fast.faster_tick(self.ns)
        item, findings = self.themed(reg)
        self.assertEqual((item["status"], len(findings)), ("retry", 1))
        state, detail = await general.inspect_one(self.ns, reg)
        self.assertEqual((state, detail["new_findings"], detail["documents_read"]), ("completed", 1, 1))
        self.assertTrue(detail["source_incomplete"])

    def test_source_snapshot_ignores_receipt_operations_and_does_not_read_original_bytes(self):
        reg = self.receipt_purchase()
        with self.connect() as con:
            before = sources.source_version(con, reg)
        path = next(iter(self.contents))
        receipt_path = path.parent / "receipt.json"
        receipt = json.loads(receipt_path.read_text())
        receipt.update(checked_at="later", state="checked", extraction_ms=99, error="temporary")
        receipt_path.write_text(json.dumps(receipt))
        with self.connect() as con, patch.object(Path, "read_bytes", side_effect=AssertionError("Read binary source")):
            self.assertEqual(sources.source_version(con, reg), before)
        original_scan = sources.original_sources
        def scan_outside_reader(number):
            with self.connect() as writer:
                writer.execute("UPDATE purchases SET last_seen='concurrent writer' WHERE reg_number=?", (reg,))
                writer.commit()
            return original_scan(number)
        with self.connect() as con, patch.object(sources, "original_sources", side_effect=scan_outside_reader):
            self.assertFalse(con.in_transaction)
            self.assertEqual(sources.source_version(con, reg), before)


if __name__ == "__main__":
    unittest.main()
