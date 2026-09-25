import os,re,json,sqlite3,asyncio,io,html,zipfile,subprocess,tempfile
from pathlib import Path
from datetime import datetime,timezone,timedelta
from urllib.parse import urlparse,urljoin,unquote,quote_plus,parse_qs
import httpx
from fastapi import FastAPI,HTTPException,Query
from fastapi.responses import HTMLResponse,FileResponse
try:
 import fitz
except Exception: fitz=None
try:
 from docx import Document
 from docx.enum.text import WD_COLOR_INDEX
except Exception: Document=None;WD_COLOR_INDEX=None
try:
 import openpyxl
except Exception: openpyxl=None
try:
 import xlrd
except Exception: xlrd=None
try:
 import pytesseract
 from PIL import Image
except Exception: pytesseract=None;Image=None
try:
 import holidays
except Exception: holidays=None
try:
 from bs4 import BeautifulSoup
except Exception: BeautifulSoup=None
try:
 from striprtf.striprtf import rtf_to_text
except Exception: rtf_to_text=None
try:
 from curl_cffi import requests as curl_requests
except Exception: curl_requests=None

APP_VERSION='1.4.0'
BASE=os.getenv('GOSPLAN_BASE','https://v2test.gosplan.info').rstrip('/')
API_KEY=os.getenv('GOSPLAN_API_KEY','').strip(); API_HEADER=os.getenv('GOSPLAN_API_HEADER','X-API-Key')
DATA=Path(os.getenv('EIS223_DATA_DIR','/data')); DATA.mkdir(parents=True,exist_ok=True)
DB=DATA/'eis223.db'; FILES=DATA/'files'; REPORTS=DATA/'reports'; FILES.mkdir(exist_ok=True); REPORTS.mkdir(exist_ok=True)
SCAN_INTERVAL=max(300,int(os.getenv('SCAN_INTERVAL_SECONDS','300')))
DETAIL_LIMIT=max(1,min(50,int(os.getenv('GOSPLAN_DETAILS_PER_CYCLE','10'))))
MAX_DOCS=max(1,min(500,int(os.getenv('MAX_DOCS_PER_DETAIL','500'))))
MAX_DOWNLOADS=max(1,min(500,int(os.getenv('MAX_ATTACHMENT_DOWNLOADS','500'))))
MAX_FILE_MB=max(1,min(100,int(os.getenv('MAX_FILE_MB','30'))))
PAGE_SIZE=max(20,min(500,int(os.getenv('GOSPLAN_PAGE_SIZE','100'))))
LAW32='https://www.consultant.ru/document/cons_doc_LAW_116964/05cd0add21b39d478c8b8af91fb2f2cd80d4a6e8/'
LAW4='https://www.consultant.ru/document/cons_doc_LAW_116964/441d00be62e3224cdc0514cffaf2a26b5b40a1c7/'
LAW3='https://www.consultant.ru/document/cons_doc_LAW_116964/fddec0f5c16a67f6fca41f9e31dfb0dcc72cc49a/'
app=FastAPI(title='ЕИС 223-ФЗ Аудитор',version=APP_VERSION)
state={'running':False,'last_start':None,'last_finish':None,'last_error':None,'last_result':None,'backfill_skip':500,'current_reg':None}
lock=asyncio.Lock()

# Наиболее распространённые организационно-правовые формы заказчиков/учреждений.
# Фильтр в UI дополнительно показывает любые новые формы, уже встреченные в БД.
ORG=[
 ('ПАО',r'\bПАО\b|публичн\w+\s+акционерн'),('НАО',r'\bНАО\b|непубличн\w+\s+акционерн'),('АО',r'\bАО\b|акционерн\w+\s+обществ'),('ООО',r'\bООО\b|ограниченн\w+\s+ответственност'),
 ('ФГУП',r'\bФГУП\b|федеральн\w+\s+государственн\w+\s+унитарн'),('ГУП',r'\bГУП\b|государственн\w+\s+унитарн'),('МУП',r'\bМУП\b|муниципальн\w+\s+унитарн'),('ФКП',r'\bФКП\b|федеральн\w+\s+казенн\w+\s+предприят'),
 ('ФГБУ',r'\bФГБУ\b|федеральн\w+\s+государственн\w+\s+бюджетн\w+\s+учрежден'),('ГБУ',r'\bГБУ\b|государственн\w+\s+бюджетн\w+\s+учрежден'),('МБУ',r'\bМБУ\b|муниципальн\w+\s+бюджетн\w+\s+учрежден'),('ФБУ',r'\bФБУ\b|федеральн\w+\s+бюджетн\w+\s+учрежден'),
 ('ФГАУ',r'\bФГАУ\b|федеральн\w+\s+государственн\w+\s+автономн\w+\s+учрежден'),('ГАУ',r'\bГАУ\b|государственн\w+\s+автономн\w+\s+учрежден'),('МАУ',r'\bМАУ\b|муниципальн\w+\s+автономн\w+\s+учрежден'),('ФАУ',r'\bФАУ\b|федеральн\w+\s+автономн\w+\s+учрежден'),
 ('ФКУ',r'\bФКУ\b|федеральн\w+\s+казенн\w+\s+учрежден'),('ГКУ',r'\bГКУ\b|государственн\w+\s+казенн\w+\s+учрежден'),('МКУ',r'\bМКУ\b|муниципальн\w+\s+казенн\w+\s+учрежден'),
 ('АНО',r'\bАНО\b|автономн\w+\s+некоммерческ'),('НКО',r'\bНКО\b|некоммерческ\w+\s+организац'),('ГК',r'\bГК\b|государственн\w+\s+корпорац'),('Фонд',r'\bфонд\b'),('Учреждение',r'\bучреждени[ея]\b')]

EXT_RE=re.compile(r'\.(pdf|docx?|xlsx?|xlsm|rtf|txt|csv|xml|zip|png|jpe?g)(?:$|[?#])',re.I)

def now(): return datetime.now(timezone.utc).isoformat()
def clean(x): return re.sub(r'\s+',' ',str(x or '')).strip()
def org_form(s):
 for n,p in ORG:
  if re.search(p,s or '',re.I): return n
 return 'Не определена'
def doc_type(name,desc=''):
 s=(str(name)+' '+str(desc)).lower().replace('ё','е')
 if 'запрос' in s and 'разъяснен' in s:return 'clarification_request'
 if 'разъяснен' in s or 'ответ на запрос' in s:return 'clarification'
 if 'отмен' in s or 'отказ от провед' in s:return 'cancellation'
 if 'протокол' in s:return 'final_protocol' if any(x in s for x in ('итог','подведен','результат')) else 'protocol'
 if 'извещен' in s:return 'notice'
 if 'документац' in s:return 'documentation'
 return 'other'
def conn():
 c=sqlite3.connect(DB,timeout=30,check_same_thread=False);c.row_factory=sqlite3.Row;return c

def init_db():
 c=conn();c.executescript('''
 CREATE TABLE IF NOT EXISTS purchases(reg_number TEXT PRIMARY KEY,title TEXT,customer TEXT,org_form TEXT,method TEXT,price TEXT,published_at TEXT,deadline TEXT,url TEXT,source TEXT,last_seen TEXT,clarifications INTEGER DEFAULT 0,protocols INTEGER DEFAULT 0,cancellations INTEGER DEFAULT 0,violations INTEGER DEFAULT 0,risks INTEGER DEFAULT 0);
 CREATE TABLE IF NOT EXISTS docs(id INTEGER PRIMARY KEY AUTOINCREMENT,reg_number TEXT,name TEXT,url TEXT,doc_type TEXT,published_at TEXT,local_path TEXT,annotated_path TEXT,status TEXT,error TEXT,metadata_json TEXT);
 CREATE TABLE IF NOT EXISTS findings(id INTEGER PRIMARY KEY AUTOINCREMENT,reg_number TEXT,code TEXT,kind TEXT,title TEXT,law TEXT,law_url TEXT,evidence TEXT,note TEXT,document TEXT,source_url TEXT);
 CREATE TABLE IF NOT EXISTS checks(id INTEGER PRIMARY KEY AUTOINCREMENT,reg_number TEXT,scope TEXT,item TEXT,status TEXT,evidence TEXT,source TEXT,source_url TEXT,document TEXT,field_path TEXT,note TEXT);
 CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY,value TEXT);
 CREATE INDEX IF NOT EXISTS idx_org ON purchases(org_form);CREATE INDEX IF NOT EXISTS idx_find_reg ON findings(reg_number);CREATE INDEX IF NOT EXISTS idx_docs_reg ON docs(reg_number);CREATE INDEX IF NOT EXISTS idx_checks_reg ON checks(reg_number);
 ''')
 cols={r['name'] for r in c.execute('PRAGMA table_info(purchases)')}
 for name,typ in [('submission_start','TEXT'),('submission_end','TEXT'),('price_num','REAL'),('card_fields_json','TEXT'),('card_source_status','TEXT')]:
  if name not in cols:c.execute(f'ALTER TABLE purchases ADD COLUMN {name} {typ}')
 c.execute('CREATE INDEX IF NOT EXISTS idx_price_num ON purchases(price_num)')
 dcols={r['name'] for r in c.execute('PRAGMA table_info(docs)')}
 for name,typ in [('text_chars','INTEGER DEFAULT 0'),('final_url','TEXT'),('checked_at','TEXT')]:
  if name not in dcols:c.execute(f'ALTER TABLE docs ADD COLUMN {name} {typ}')
 fcols={r['name'] for r in c.execute('PRAGMA table_info(findings)')}
 for name,typ in [('location','TEXT'),('match_text','TEXT')]:
  if name not in fcols:c.execute(f'ALTER TABLE findings ADD COLUMN {name} {typ}')
 ccols={r['name'] for r in c.execute('PRAGMA table_info(checks)')}
 for name,typ in [('law','TEXT'),('law_url','TEXT'),('location','TEXT')]:
  if name not in ccols:c.execute(f'ALTER TABLE checks ADD COLUMN {name} {typ}')
 pcols={r['name'] for r in c.execute('PRAGMA table_info(purchases)')}
 for name,typ in [('documents_found','INTEGER DEFAULT 0'),('documents_checked','INTEGER DEFAULT 0'),('checks_ok','INTEGER DEFAULT 0'),('checks_attention','INTEGER DEFAULT 0')]:
  if name not in pcols:c.execute(f'ALTER TABLE purchases ADD COLUMN {name} {typ}')
 c.commit();c.close()
init_db()

def walk(x):
 if isinstance(x,dict):
  yield x
  for v in x.values():yield from walk(v)
 elif isinstance(x,list):
  for v in x:yield from walk(v)
def getv(obj,*keys):
 ks={k.lower() for k in keys}
 for d in walk(obj):
  for k,v in d.items():
   if str(k).lower() in ks and v not in (None,'',[]):return v
 return None
def txt(v):
 if v is None:return ''
 if isinstance(v,(str,int,float)):return clean(v)
 if isinstance(v,dict):
  for k in ('fullName','name','shortName','organizationName','customerName','value','text','title'):
   if v.get(k):return txt(v[k])
 if isinstance(v,list):
  for x in v:
   q=txt(x)
   if q:return q
 return ''
def regnum(obj):
 s=txt(getv(obj,'purchaseNumber','purchase_number','registryNumber','regNumber','noticeNumber','number'))
 m=re.search(r'\d{10,}',s);return m.group(0) if m else s
def old_eis_url(reg):return f'https://zakupki.gov.ru/223/purchase/public/purchase/info/common-info.html?regNumber={reg}'
def old_eis_docs_url(reg):return f'https://zakupki.gov.ru/223/purchase/public/purchase/info/documents.html?regNumber={reg}'
def eis_search_url(reg):
 return ('https://zakupki.gov.ru/epz/order/extendedsearch/results.html?searchString='+quote_plus(str(reg))+
         '&morphology=on&search-filter=%D0%94%D0%B0%D1%82%D0%B5+%D1%80%D0%B0%D0%B7%D0%BC%D0%B5%D1%89%D0%B5%D0%BD%D0%B8%D1%8F&pageNumber=1&sortDirection=false&recordsPerPage=_10&showLotsInfoHidden=false&sortBy=UPDATE_DATE&fz223=on&af=on&currencyIdGeneral=-1')
def current_223_common(notice_info_id):return f'https://zakupki.gov.ru/epz/order/notice/notice223/common-info.html?noticeInfoId={notice_info_id}'
def current_223_pages(common_url):
 if not common_url:return {}
 return {
  'common':common_url,
  'lots':common_url.replace('common-info.html','lot-list.html'),
  'documents':common_url.replace('common-info.html','documents.html'),
 }
def eis_url(reg):return old_eis_url(reg)
def eis_docs_url(reg):return old_eis_docs_url(reg)
def eis_page_urls(reg):
 base='https://zakupki.gov.ru/223/purchase/public/purchase/info/'
 return {'common':base+f'common-info.html?regNumber={reg}','lots':base+f'lot-list.html?regNumber={reg}','documents':base+f'documents.html?regNumber={reg}','changes':base+f'change-history.html?regNumber={reg}'}

def dt(v):
 if not v:return None
 s=clean(v)
 # ЕИС часто добавляет «(МСК)»/«МСК+N»; для проверки порядка дат берём календарную дату/время.
 m=re.search(r'(\d{2}\.\d{2}\.\d{4})(?:\s+(\d{1,2}:\d{2})(?::(\d{2}))?)?',s)
 if m:
  raw=m.group(1)+((' '+m.group(2)) if m.group(2) else '')+((':'+m.group(3)) if m.group(3) else '')
  for f in ('%d.%m.%Y %H:%M:%S','%d.%m.%Y %H:%M','%d.%m.%Y'):
   try:return datetime.strptime(raw,f)
   except:pass
 for z in (s,s.replace('Z','+00:00')):
  try:return datetime.fromisoformat(z)
  except:pass
 for f in ('%Y-%m-%d %H:%M:%S','%Y-%m-%d'):
  try:return datetime.strptime(s[:19],f)
  except:pass
 return None
def wdays(a,b):
 if not a or not b or b<=a:return 0
 hs=None
 if holidays:
  try:hs=holidays.RU(years=range(a.year,b.year+1))
  except:pass
 n=0;d=a.date()+timedelta(days=1)
 while d<=b.date():
  if d.weekday()<5 and (hs is None or d not in hs):n+=1
  d+=timedelta(days=1)
 return n

def price_num(v):
 if v is None:return None
 if isinstance(v,(int,float)):return float(v)
 s=clean(v).replace('\xa0',' ')
 # Ищем денежную величину, а не склеиваем все цифры из длинного текста.
 vals=re.findall(r'(?<!\d)(\d{1,3}(?:[ \u00a0]\d{3})+(?:[,.]\d{1,2})?|\d+(?:[,.]\d{1,2})?)(?!\d)',s)
 if not vals:return None
 best=None
 for raw in vals:
  z=raw.replace(' ','').replace('\u00a0','').replace(',','.')
  try:
   x=float(z)
   if best is None or x>best:best=x
  except:pass
 return best

def headers(json_mode=True,referer=None):
 h={'User-Agent':'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/140 Safari/537.36','Accept-Language':'ru-RU,ru;q=0.9,en;q=0.5','Cache-Control':'no-cache'}
 h['Accept']='application/json' if json_mode else 'text/html,application/xhtml+xml,application/xml;q=0.9,application/pdf,application/octet-stream,*/*;q=0.8'
 if referer:h['Referer']=referer
 if API_KEY and json_mode:h[API_HEADER]=API_KEY
 return h
async def jget(client,path,params=None):
 last=None
 for _ in range(3):
  r=await client.get(BASE+path,params=params,headers=headers(True))
  if r.status_code!=429:r.raise_for_status();return r.json()
  last=r;wait=60
  for k in ('retry-after','ratelimit-reset'):
   try:
    if r.headers.get(k):wait=max(2,min(180,int(float(r.headers[k]))+2));break
   except:pass
  await asyncio.sleep(wait)
 if last:last.raise_for_status()

def flatten_scalars(obj,path='',out=None,limit=2500):
 if out is None:out=[]
 if len(out)>=limit:return out
 if isinstance(obj,dict):
  for k,v in obj.items():
   p=f'{path}.{k}' if path else str(k);flatten_scalars(v,p,out,limit)
 elif isinstance(obj,list):
  for i,v in enumerate(obj[:200]):flatten_scalars(v,f'{path}[{i}]',out,limit)
 elif obj not in (None,''):
  s=clean(obj)
  if s and len(s)<=4000:out.append((path,s))
 return out

def _looks_like_name(v):
 s=clean(v)
 if not s or len(s)<3:return False
 if re.fullmatch(r'[+\\d .()/-]+',s):return False
 if re.fullmatch(r'\\d{8,15}',re.sub(r'\\D','',s)) and not re.search(r'[A-Za-zА-Яа-яЁё]',s):return False
 return bool(re.search(r'[A-Za-zА-Яа-яЁё]',s))

def api_customer(detail):
 # Сначала точные поля с названием заказчика, затем вложенные customer/organization.
 for key in ('customerName','customerFullName','customerOrganizationName','organizationName','placerName','customerShortName'):
  v=txt(getv(detail,key))
  if _looks_like_name(v):return v
 for key in ('customer','customerInfo','organization','placer'):
  obj=getv(detail,key);v=txt(obj)
  if _looks_like_name(v):return v
 fields=flatten_scalars(detail,limit=5000)
 scored=[]
 for path,val in fields:
  lp=path.lower()
  if any(x in lp for x in ('inn','kpp','ogrn','tax','id','code')):continue
  if not _looks_like_name(val):continue
  score=0
  if 'customer' in lp:score+=30
  if 'placer' in lp:score+=25
  if 'organization' in lp:score+=20
  if lp.endswith(('fullname','organizationname','customername','shortname','.name')):score+=20
  if score:scored.append((score,-len(path),val))
 return max(scored)[2] if scored else ''

def api_price(detail):
 keys=('initialSum','initialPrice','maxPrice','price','lotPrice','initialContractPrice','initialMaxPrice','initialMaximumPrice','nmcd','nmcp','startPrice','initialContractSum')
 for key in keys:
  v=txt(getv(detail,key));n=price_num(v)
  if n is not None and n>=0:return v
 best=[]
 for path,val in flatten_scalars(detail,limit=5000):
  lp=path.lower()
  if any(x in lp for x in ('price','sum','nmcd','nmcp','initial','maximum','maxprice')) and not any(x in lp for x in ('id','currency','code')):
   n=price_num(val)
   if n is not None and n>=0:best.append((n,val,path))
 return max(best,key=lambda x:x[0])[1] if best else ''

def submission_dates(detail):
 start=txt(getv(detail,'submissionStartDate','submissionStartDateTime','applicationStartDate','applicationStartDateTime','submissionOpenDate','applicationOpenDate','startSubmissionDate','applicationSubmissionStartDate'))
 end=txt(getv(detail,'submissionCloseDate','submissionEndDate','submissionEndDateTime','applicationEndDate','applicationEndDateTime','submissionDeadline','applicationDeadline','endSubmissionDate','applicationSubmissionEndDate'))
 return start,end

def extract_docs(detail):
 out=[];seen=set()
 for d in walk(detail):
  name=txt(d.get('fileName') or d.get('filename') or d.get('documentName') or d.get('docName') or d.get('name'))
  desc=txt(d.get('description') or d.get('documentType') or d.get('typeName') or d.get('title'))
  urls=[]
  for k in ('url','downloadUrl','fileUrl','href','link','sourceUrl','attachmentUrl','downloadLink'):
   u=txt(d.get(k))
   if u and u.startswith(('http://','https://')) and u not in urls:urls.append(u)
  url=urls[0] if urls else ''
  label=(name+' '+desc).strip(); low=label.lower()
  if not (url or name):continue
  if not (EXT_RE.search(url) or EXT_RE.search(name) or any(x in low for x in ('протокол','разъяснен','отмен','документац','извещен','запрос','техническ','обоснован','форма заявки'))):continue
  key=(name,url)
  if key in seen:continue
  seen.add(key);meta={}
  for k,v in d.items():
   lk=str(k).lower()
   if any(x in lk for x in ('date','time','publish','sign','decision','request','create','id','guid')) and not isinstance(v,(dict,list)):meta[str(k)]=v
  if len(urls)>1:meta['alt_urls']=urls[1:]
  out.append({'name':name or desc or 'Документ','desc':desc,'url':url,'type':doc_type(name,desc),'published':txt(d.get('publishDate') or d.get('publicationDate') or d.get('publishedAt') or d.get('createDate')),'meta':meta,'text':'','local':None,'annotated':None,'status':'metadata','error':None,'final_url':None})
 return out[:MAX_DOCS]

def safe_host(url):
 try:
  p=urlparse(url);h=(p.hostname or '').lower()
  return p.scheme in ('http','https') and any(h==x or h.endswith('.'+x) for x in ('zakupki.gov.ru','gosplan.info','roskazna.gov.ru'))
 except:return False
def safe_name(s):return re.sub(r'[^0-9A-Za-zА-Яа-яЁё._-]+','_',unquote(s or 'file'))[:180]

def _clean_tooltip(v):
 if not v:return ''
 try:return clean(BeautifulSoup(html.unescape(v),'html.parser').get_text(' ',strip=True)) if BeautifulSoup else clean(re.sub(r'<[^>]+>',' ',html.unescape(v)))
 except:return clean(v)

def parse_card_html(raw):
 pairs=[];visible=''
 try:
  if BeautifulSoup:
   soup=BeautifulSoup(raw or '','html.parser')
   # Ключевые блоки актуальной карточки 223-ФЗ.
   for title_cls,value_cls in (('registry-entry__body-title','registry-entry__body-value'),('price-block__title','price-block__value'),('data-block__title','data-block__value'),('common-text__title','common-text__value'),('section__title','section__info'),('attachment__text','attachment__value')):
    for t in soup.select('.'+title_cls):
     label=clean(t.get_text(' ',strip=True));v=t.find_next_sibling(class_=value_cls)
     if v is None and getattr(t,'parent',None):v=t.parent.select_one('.'+value_cls)
     value=clean(v.get_text(' ',strip=True)) if v else ''
     if label and value:pairs.append((label,value))
   for tr in soup.find_all('tr'):
    cells=[clean(c.get_text(' ',strip=True)) for c in tr.find_all(['th','td'],recursive=False)]
    if len(cells)>=2 and cells[0] and any(cells[1:]):pairs.append((cells[0],' | '.join(x for x in cells[1:] if x)))
   for dtel in soup.find_all('dt'):
    dd=dtel.find_next_sibling('dd')
    if dd:pairs.append((clean(dtel.get_text(' ',strip=True)),clean(dd.get_text(' ',strip=True))))
   work=BeautifulSoup(raw or '','html.parser')
   for x in work(['script','style','noscript']):x.decompose()
   visible='\n'.join(clean(x) for x in work.stripped_strings if clean(x))
  else:visible=clean(re.sub(r'<[^>]+>',' ',raw or ''))
 except Exception:pass
 seen=set();uniq=[]
 for k,v in pairs:
  k,v=clean(k)[:350],clean(v)[:3000]
  if not k or not v:continue
  sig=(k.lower(),v.lower())
  if sig in seen:continue
  seen.add(sig);uniq.append((k,v))
 return uniq[:2500],visible[:300000]

def eis_blocked(raw,status_code=200):
 s=(raw or '').lower()
 markers=('access denied','request rejected','captcha','robot check','доступ ограничен','проверка браузера','cloudflare','security check')
 return status_code in (401,403,429) or any(x in s for x in markers)

async def eis_request(client,url,referer='https://zakupki.gov.ru/',max_bytes=8*1024*1024):
 errors=[]
 try:
  r=await client.get(url,headers=headers(False,referer),follow_redirects=True)
  data=r.content[:max_bytes+1];raw=r.text if len(r.content)<=max_bytes else r.content[:max_bytes].decode(r.encoding or 'utf-8','ignore')
  if r.status_code<400 and not eis_blocked(raw,r.status_code):
   return {'status_code':r.status_code,'url':str(r.url),'headers':dict(r.headers),'content':data,'text':raw,'engine':'httpx'}
  errors.append(f'httpx HTTP {r.status_code}')
 except Exception as e:errors.append('httpx '+repr(e))
 if curl_requests:
  def _curl():
   rr=curl_requests.get(url,headers=headers(False,referer),impersonate='chrome',timeout=30,allow_redirects=True)
   content=bytes(rr.content[:max_bytes+1]);enc=getattr(rr,'encoding',None) or 'utf-8'
   try:text=content[:max_bytes].decode(enc,'ignore')
   except:text=content[:max_bytes].decode('utf-8','ignore')
   return rr,content,text
  try:
   rr,data,raw=await asyncio.to_thread(_curl)
   if rr.status_code<400 and not eis_blocked(raw,rr.status_code):
    return {'status_code':rr.status_code,'url':str(rr.url),'headers':dict(rr.headers),'content':data,'text':raw,'engine':'curl_cffi'}
   errors.append(f'curl_cffi HTTP {rr.status_code}')
  except Exception as e:errors.append('curl_cffi '+repr(e))
 raise RuntimeError('; '.join(errors)[:1200])

def _current_common_from_detail(detail):
 for path,val in flatten_scalars(detail,limit=5000):
  s=clean(val)
  if 'zakupki.gov.ru' in s and '/epz/order/notice/notice223/' in s and 'common-info' in s:
   m=re.search(r'https?://[^\s"<>]+',s);return html.unescape(m.group(0)) if m else s
  if 'noticeinfoid' in path.lower():
   m=re.search(r'\d+',s)
   if m:return current_223_common(m.group(0))
 blob=json.dumps(detail,ensure_ascii=False,default=str)
 m=re.search(r'noticeInfoId(?:%3D|=|["\s:]+)(\d+)',blob,re.I)
 return current_223_common(m.group(1)) if m else ''

async def resolve_eis_common(client,reg,detail=None):
 direct=_current_common_from_detail(detail or {})
 if direct:return direct,{'method':'api-field'}
 # Актуальная выдача ЕИС содержит noticeInfoId, которого нет в реестровом номере.
 search=eis_search_url(reg)
 try:
  q=await eis_request(client,search,'https://zakupki.gov.ru/')
  if BeautifulSoup:
   soup=BeautifulSoup(q['text'],'html.parser')
   anchors=soup.select('a[href*="/epz/order/notice/notice223/"][href*="common-info"]')
   chosen=None
   for a in anchors:
    if str(reg) in clean(a.get_text(' ',strip=True)):chosen=a;break
   if chosen is None and len(anchors)==1:chosen=anchors[0]
   if chosen and chosen.get('href'):
    return urljoin(q['url'],chosen.get('href')),{'method':'eis-search','search_url':q['url'],'engine':q.get('engine')}
 except Exception as e:search_error=str(e)
 # Последний fallback: старый публичный URL; если ЕИС редиректит, ниже зафиксируем конечный актуальный URL.
 return old_eis_url(reg),{'method':'legacy-fallback','search_error':locals().get('search_error')}

def _tab_urls_from_html(raw,base_url):
 out={}
 if BeautifulSoup:
  try:
   soup=BeautifulSoup(raw or '','html.parser')
   for a in soup.select('.tabsNav a[href], a[href]'):
    text=clean(a.get_text(' ',strip=True)).lower().replace('ё','е');href=a.get('href')
    if not href:continue
    u=urljoin(base_url,href)
    if 'общ' in text and 'информац' in text:out.setdefault('common',u)
    elif 'лот' in text:out.setdefault('lots',u)
    elif 'документ' in text:out.setdefault('documents',u)
    elif 'измен' in text or 'журнал' in text:out.setdefault('changes',u)
  except:pass
 return out

def _pair_value(pairs,*labels):
 norms=[clean(x).lower().replace('ё','е') for x in labels]
 for k,v in pairs:
  nk=clean(k).lower().replace('ё','е')
  if any(nk==x or x in nk for x in norms):return clean(v)
 return ''

def extract_eis_summary(card,reg):
 pairs=card.get('pairs',[]);common=(card.get('pages') or {}).get('common',{});lots=(card.get('pages') or {}).get('lots',{})
 raw=common.get('html') or '';soup=BeautifulSoup(raw,'html.parser') if BeautifulSoup and raw else None
 def labeled(label):
  if not soup:return ''
  nl=clean(label)
  for tc,vc in (('registry-entry__body-title','registry-entry__body-value'),('price-block__title','price-block__value'),('data-block__title','data-block__value'),('common-text__title','common-text__value')):
   for t in soup.select('.'+tc):
    if clean(t.get_text(' ',strip=True))==nl:
     v=t.find_next_sibling(class_=vc)
     if v is None and getattr(t,'parent',None):v=t.parent.select_one('.'+vc)
     if v:return clean(v.get_text(' ',strip=True))
  return ''
 customer=labeled('Заказчик') or _pair_value(pairs,'Наименование организации','Заказчик','Организация, осуществляющая размещение')
 title=labeled('Объект закупки') or _pair_value(pairs,'Наименование закупки','Наименование объекта закупки','Объект закупки')
 price=labeled('Начальная цена') or labeled('Начальная (максимальная) цена договора') or _pair_value(pairs,'Начальная (максимальная) цена договора','Начальная цена','НМЦД','Сведения о цене договора')
 if not price:
  lt=(lots.get('visible_text') or '')+'\n'+(lots.get('html') or '')
  m=re.search(r'(?:Начальная\s*\(максимальная\)\s*цена\s*договора|Начальная\s+цена|НМЦД)[^0-9]{0,80}([0-9][0-9 \u00a0]*(?:[,.][0-9]{1,2})?\s*(?:₽|руб\.?|Российский рубль)?)',lt,re.I)
  if m:price=clean(m.group(1))
 method=_pair_value(pairs,'Способ осуществления закупки','Способ закупки','Способ проведения закупки')
 published=labeled('Размещено') or _pair_value(pairs,'Дата размещения извещения','Дата размещения')
 start=_pair_value(pairs,'Дата начала срока подачи заявок','Дата и время начала срока подачи заявок','Начало подачи заявок')
 end=_pair_value(pairs,'Дата и время окончания срока подачи заявок (по местному времени заказчика)','Дата и время окончания срока подачи заявок','Дата окончания срока подачи заявок','Окончание подачи заявок') or labeled('Окончание подачи заявок')
 return {'reg':str(reg),'customer':customer,'title':title,'price':price,'price_num':price_num(price),'method':method,'published':published,'submission_start':start,'submission_end':end}

async def fetch_eis_card(client,reg,detail=None):
 result={'status':'error','pairs':[],'visible_text':'','html':'','error':None,'pages':{},'summary':{},'discovery':{}}
 errors=[];all_pairs=[];all_visible=[];success=0
 common_url,disc=await resolve_eis_common(client,reg,detail);result['discovery']=disc
 # Основную страницу загружаем первой и по ней определяем актуальные URL вкладок.
 first={'url':common_url,'status':'error','http_status':None,'final_url':None,'pairs':[],'visible_text':'','html':'','error':None}
 try:
  rr=await eis_request(client,common_url,'https://zakupki.gov.ru/')
  pairs,visible=parse_card_html(rr['text']);first.update(status='ok',http_status=rr['status_code'],final_url=rr['url'],pairs=pairs,visible_text=visible,html=rr['text'][:800000],engine=rr.get('engine'))
  success+=1;all_pairs.extend(('common: '+a,b) for a,b in pairs);all_visible.append('[common]\n'+visible)
 except Exception as e:first['error']=str(e)[:1000];errors.append('common: '+str(e))
 result['pages']['common']=first
 tabs=_tab_urls_from_html(first.get('html',''),first.get('final_url') or common_url)
 if '/epz/order/notice/notice223/' in (first.get('final_url') or common_url):
  tabs={**current_223_pages(first.get('final_url') or common_url),**tabs}
 else:
  tabs={**eis_page_urls(reg),**tabs}
 for key in ('lots','documents','changes'):
  url=tabs.get(key)
  if not url:continue
  pg={'url':url,'status':'error','http_status':None,'final_url':None,'pairs':[],'visible_text':'','html':'','error':None}
  try:
   rr=await eis_request(client,url,first.get('final_url') or common_url)
   pairs,visible=parse_card_html(rr['text']);pg.update(status='ok',http_status=rr['status_code'],final_url=rr['url'],pairs=pairs,visible_text=visible,html=rr['text'][:1000000],engine=rr.get('engine'))
   success+=1;all_pairs.extend((key+': '+a,b) for a,b in pairs);all_visible.append('['+key+']\n'+visible)
  except Exception as e:pg['error']=str(e)[:1000];errors.append(key+': '+str(e))
  result['pages'][key]=pg
 seen=set();pairs=[]
 for k,v in all_pairs:
  sig=(clean(k).lower(),clean(v).lower())
  if sig in seen:continue
  seen.add(sig);pairs.append((k,v))
 result['pairs']=pairs[:4000];result['visible_text']='\n'.join(all_visible)[:500000];result['html']='\n'.join(p.get('html','') for p in result['pages'].values())[:1500000]
 common_ok=result['pages'].get('common',{}).get('status')=='ok';docs_ok=result['pages'].get('documents',{}).get('status')=='ok'
 result['status']='ok' if common_ok and docs_ok else ('partial' if success else 'error');result['error']='; '.join(errors)[:2000] if errors else None
 result['summary']=extract_eis_summary(result,reg)
 return result

def _candidate_urls_from_html(raw,base_url):
 found=[]
 def add(u):
  if not u:return
  u=html.unescape(str(u)).replace('\\/','/').strip()
  # В актуальном 223-ФЗ прямая ссылка выглядит /223/purchase/public/download/download.html?id=...
  mm=re.search(r"(https?://[^\s<>]+|/[^\s<>]+)",u)
  if mm:u=mm.group(1)
  full=urljoin(base_url,u);low=full.lower()
  if safe_host(full) and (EXT_RE.search(low) or '/purchase/public/download/download.html' in low or '/filestore/public/1.0/download/' in low or any(x in low for x in ('attachment','getfile','downloadfile'))):
   if full not in found:found.append(full)
 if BeautifulSoup:
  try:
   soup=BeautifulSoup(raw or '','html.parser')
   for tag in soup.find_all(True):
    for attr in ('href','data-href','data-url','data-download-url','data-file-url','onclick'):
     if tag.get(attr):add(tag.get(attr))
  except:pass
 else:
  for u in re.findall(r'https?://[^\s"<>]+',raw or '',re.I):add(u)
 return found[:1000]

def _doc_metadata(link):
 meta={}
 if not BeautifulSoup:return meta
 try:
  for ancestor in list(link.parents)[:6]:
   for tc,vc in (('attachment__text','attachment__value'),('section__attrib','section__value')):
    for t in ancestor.select('.'+tc):
     label=clean(t.get_text(' ',strip=True));v=t.find_next_sibling(class_=vc)
     if v is None and getattr(t,'parent',None):v=t.parent.select_one('.'+vc)
     value=clean(v.get_text(' ',strip=True)) if v else ''
     if label and value:meta.setdefault(label,value)
   if meta:return meta
 except:pass
 return meta

async def discover_eis_documents(client,reg,card=None):
 out=[];page=(card or {}).get('pages',{}).get('documents',{});raw=page.get('html') or '';base=page.get('final_url') or page.get('url') or old_eis_docs_url(reg)
 if not raw:
  try:
   rr=await eis_request(client,base,(card or {}).get('pages',{}).get('common',{}).get('final_url') or old_eis_url(reg));raw=rr['text'];base=rr['url']
  except Exception:return []
 seen=set()
 if BeautifulSoup:
  try:
   soup=BeautifulSoup(raw,'html.parser')
   for a in soup.select('a[href]'):
    href=a.get('href') or '';u=urljoin(base,href);low=u.lower()
    if not safe_host(u):continue
    if not ('/223/purchase/public/download/download.html' in low or '/filestore/public/1.0/download/' in low):continue
    if u in seen:continue
    seen.add(u)
    name=_clean_tooltip(a.get('data-tooltip')) or _clean_tooltip(a.get('title')) or clean(a.get_text(' ',strip=True)) or 'Документ ЕИС'
    meta=_doc_metadata(a);meta['source']='eis-documents-page'
    out.append({'name':name,'desc':'Оригинал со страницы «Документы» ЕИС','url':u,'type':doc_type(name,'Документ ЕИС'),'published':meta.get('Размещено') or meta.get('Дата размещения') or '','meta':meta,'text':'','local':None,'annotated':None,'status':'metadata','error':None,'final_url':None})
  except:pass
 # Дополнительный универсальный проход нужен для нестандартных ссылок/старых редакций.
 names={}
 if BeautifulSoup:
  try:
   soup=BeautifulSoup(raw,'html.parser')
   for a in soup.find_all('a',href=True):names[urljoin(base,a.get('href'))]=_clean_tooltip(a.get('data-tooltip')) or _clean_tooltip(a.get('title')) or clean(a.get_text(' ',strip=True))
  except:pass
 for u in _candidate_urls_from_html(raw,base):
  if u in seen:continue
  seen.add(u);name=names.get(u) or Path(unquote(urlparse(u).path)).name or 'Документ ЕИС'
  out.append({'name':name,'desc':'Оригинал со страницы «Документы» ЕИС','url':u,'type':doc_type(name,'Документ ЕИС'),'published':'','meta':{'source':'eis-documents-page'},'text':'','local':None,'annotated':None,'status':'metadata','error':None,'final_url':None})
 return out[:MAX_DOCS]

def merge_docs(a,b):
 out=[];seen=set()
 for d in a+b:
  key=(clean(d.get('name')).lower(),clean(d.get('url')).lower())
  if key in seen:continue
  seen.add(key);out.append(d)
 return out[:MAX_DOCS]

def filename_from_headers(headers_map,fallback):
 cd=(headers_map or {}).get('content-disposition','') or (headers_map or {}).get('Content-Disposition','')
 m=re.search(r"filename\*=UTF-8''([^;]+)",cd,re.I) or re.search(r'filename="?([^";]+)',cd,re.I)
 return safe_name(unquote(m.group(1)) if m else fallback)

def looks_html(data,ct):
 head=bytes(data[:1500]).lstrip().lower()
 return 'text/html' in (ct or '').lower() or head.startswith(b'<!doctype html') or head.startswith(b'<html')

async def fetch_binary(client,url,reg,referer=None):
 rr=await eis_request(client,url,referer or old_eis_url(reg),max_bytes=MAX_FILE_MB*1024*1024)
 if not safe_host(rr['url']):raise ValueError('Переадресация на недопустимый домен')
 data=rr['content'];ct=(rr['headers'].get('content-type') or rr['headers'].get('Content-Type') or '').lower()
 if len(data)>MAX_FILE_MB*1024*1024:raise ValueError(f'Файл больше {MAX_FILE_MB} МБ')
 if looks_html(data,ct):
  links=_candidate_urls_from_html(rr['text'],rr['url'])
  if links:
   rr=await eis_request(client,links[0],rr['url'],max_bytes=MAX_FILE_MB*1024*1024);data=rr['content'];ct=(rr['headers'].get('content-type') or rr['headers'].get('Content-Type') or '').lower()
 if looks_html(data,ct):raise ValueError('Ссылка вернула HTML-страницу вместо файла')
 return rr,data,ct

def _cache_docs(reg):
 try:
  c=conn();rows=[dict(r) for r in c.execute('SELECT url,final_url,local_path FROM docs WHERE reg_number=?',(reg,))];c.close()
  out={}
  for r in rows:
   if r.get('local_path') and Path(r['local_path']).exists():
    if r.get('url'):out[r['url']]=r['local_path']
    if r.get('final_url'):out[r['final_url']]=r['local_path']
  return out
 except:return {}

async def download_doc(client,reg,d,cache=None):
 urls=[d.get('url')]+list(d.get('meta',{}).get('alt_urls') or []);urls=[u for u in urls if u and safe_host(u)]
 if not urls:d['status']='error';d['error']='Нет допустимой ссылки на оригинал';return
 cache=cache or {}
 for u in urls:
  pth=cache.get(u)
  if pth and Path(pth).exists():
   d['local']=pth;d['final_url']=u;d['text']=extract_text(pth);chars=len(clean(d['text']));d['status']='checked' if chars>=3 else 'downloaded_no_text';d['error']=None if chars>=3 else 'Сохранённый оригинал есть, но текст не извлечён';return
 last=None
 for u in urls:
  try:
   rr,data,ct=await fetch_binary(client,u,reg)
   fallback=d.get('name') or Path(urlparse(rr['url']).path).name or 'file';name=filename_from_headers(rr.get('headers'),fallback)
   suffix=Path(name).suffix.lower()
   if not suffix:
    if data[:4]==b'%PDF':name+='.pdf'
    elif data[:2]==b'PK':name+='.zip'
    elif data[:4]==b'\xd0\xcf\x11\xe0':name+=('.xls' if 'excel' in ct or 'sheet' in ct else '.doc')
    elif 'word' in ct:name+='.docx'
    elif 'sheet' in ct or 'excel' in ct:name+='.xlsx'
    elif 'xml' in ct:name+='.xml'
    else:name+='.bin'
   folder=FILES/reg/'original';folder.mkdir(parents=True,exist_ok=True);p=folder/name
   # Не оставляем HTML под видом документа.
   if looks_html(data,ct):raise ValueError('Получен HTML вместо оригинала')
   p.write_bytes(data);d['local']=str(p);d['final_url']=rr['url'];d['text']=extract_text(p);chars=len(clean(d['text']))
   d['status']='checked' if chars>=3 else 'downloaded_no_text';d['error']=None if chars>=3 else 'Оригинал скачан, но содержимое извлечь не удалось'
   return
  except Exception as e:last=e
 d['status']='error';d['error']=str(last)[:700] if last else 'Не удалось скачать оригинал'

async def download_all_documents(client,reg,docs):
 cache=_cache_docs(reg);sem=asyncio.Semaphore(3)
 async def one(d):
  async with sem:await download_doc(client,reg,d,cache)
 batch=docs[:MAX_DOWNLOADS];await asyncio.gather(*(one(d) for d in batch))
 for d in docs[MAX_DOWNLOADS:]:d['status']='error';d['error']='Не обработан из-за защитного лимита загрузок'

def extract_text(p,depth=0):
 p=Path(p);ext=p.suffix.lower()
 if depth>2:return ''
 try:
  if ext=='.pdf' and fitz:
   doc=fitz.open(p);parts=[]
   for pg in doc:
    t=pg.get_text('text') or ''
    if len(t.strip())<30 and pytesseract and Image:
     try:
      pix=pg.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False);t=pytesseract.image_to_string(Image.open(io.BytesIO(pix.tobytes('png'))),lang='rus+eng')
     except:pass
    parts.append(t)
   doc.close();return '\n'.join(parts)
  if ext=='.docx' and Document:
   d=Document(p);return '\n'.join([x.text for x in d.paragraphs]+[' | '.join(c.text for c in row.cells) for t in d.tables for row in t.rows])
  if ext=='.doc':
   try:return subprocess.run(['antiword',str(p)],capture_output=True,text=True,timeout=20,errors='ignore').stdout
   except:return ''
  if ext in ('.xlsx','.xlsm') and openpyxl:
   wb=openpyxl.load_workbook(p,read_only=True,data_only=True);out=[]
   for ws in wb.worksheets:
    for row in ws.iter_rows(values_only=True):
     vals=[str(v) for v in row if v is not None]
     if vals:out.append(' | '.join(vals))
   return '\n'.join(out)
  if ext=='.xls' and xlrd:
   wb=xlrd.open_workbook(str(p));out=[]
   for ws in wb.sheets():
    for i in range(ws.nrows):
     vals=[str(x) for x in ws.row_values(i) if str(x).strip()]
     if vals:out.append(' | '.join(vals))
   return '\n'.join(out)
  if ext=='.rtf':
   raw=p.read_text(errors='ignore');return rtf_to_text(raw) if rtf_to_text else raw
  if ext in ('.txt','.csv','.xml'):
   raw=p.read_text(errors='ignore')
   return clean(re.sub(r'<[^>]+>',' ',raw)) if ext=='.xml' else raw
  if ext in ('.png','.jpg','.jpeg') and pytesseract and Image:return pytesseract.image_to_string(Image.open(p),lang='rus+eng')
  if ext=='.zip':
   parts=[];total=0
   with zipfile.ZipFile(p) as z:
    for info in z.infolist()[:50]:
     if info.is_dir() or info.file_size>MAX_FILE_MB*1024*1024:continue
     total+=info.file_size
     if total>100*1024*1024:break
     member=Path(info.filename)
     if member.suffix.lower() not in ('.pdf','.doc','.docx','.xls','.xlsx','.xlsm','.rtf','.txt','.csv','.xml','.png','.jpg','.jpeg','.zip'):continue
     data=z.read(info)
     with tempfile.NamedTemporaryFile(suffix=member.suffix,delete=False) as tf:tf.write(data);tmp=tf.name
     try:
      t=extract_text(tmp,depth+1)
      if t:parts.append(f'--- {member.name} ---\n{t}')
     finally:
      try:os.unlink(tmp)
      except:pass
   return '\n'.join(parts)
 except:return ''
 return ''

def mdate(d,*need):
 for k,v in d.get('meta',{}).items():
  lk=k.lower()
  if all(n in lk for n in need):
   z=dt(v)
   if z:return z
 return dt(d.get('published')) if ('publish' in need or 'publication' in need) else None

def compact_seed(s):
 s=clean(s)
 if not s:return ''
 # Сначала берём фрагмент без переносов; слишком длинные контексты плохо ищутся в PDF/DOCX.
 return s[:120]

def locate_in_file(path,needle):
 if not path or not needle:return ''
 p=Path(path);ext=p.suffix.lower();needle=clean(needle)
 try:
  if ext=='.pdf' and fitz:
   doc=fitz.open(p)
   for i,pg in enumerate(doc):
    text=pg.get_text('text') or ''
    if needle.lower() in clean(text).lower() or (len(needle)>25 and needle[:50].lower() in clean(text).lower()):doc.close();return f'страница {i+1}'
   doc.close()
  elif ext=='.docx' and Document:
   d=Document(p)
   for i,par in enumerate(d.paragraphs,1):
    if needle.lower() in clean(par.text).lower() or (len(needle)>25 and needle[:50].lower() in clean(par.text).lower()):return f'абзац {i}'
   for ti,t in enumerate(d.tables,1):
    for ri,row in enumerate(t.rows,1):
     for ci,cell in enumerate(row.cells,1):
      if needle.lower() in clean(cell.text).lower() or (len(needle)>25 and needle[:50].lower() in clean(cell.text).lower()):return f'таблица {ti}, строка {ri}, ячейка {ci}'
  elif ext in ('.xlsx','.xlsm') and openpyxl:
   wb=openpyxl.load_workbook(p,read_only=True,data_only=True)
   for ws in wb.worksheets:
    for row in ws.iter_rows():
     for cell in row:
      val=clean(cell.value)
      if needle.lower() in val.lower() or (len(needle)>25 and needle[:50].lower() in val.lower()):return f'лист «{ws.title}», ячейка {cell.coordinate}'
  else:
   t=extract_text(p);pos=t.lower().find(needle.lower())
   if pos>=0:return f'текст документа, позиция ~{pos+1}'
 except Exception:pass
 return ''

def finding(code,kind,title,law,law_url,evidence,note,doc=None,match_text='',location=''):
 if doc and not location and match_text:location=locate_in_file(doc.get('local'),match_text)
 return {'code':code,'kind':kind,'title':title,'law':law,'law_url':law_url,'evidence':clean(evidence)[:1800],'note':note,'document':doc['name'] if doc else None,'source_url':doc['url'] if doc else None,'location':location,'match_text':clean(match_text)[:500]}
def check(scope,item,status,evidence='',source='',source_url='',document=None,field_path='',note='',law='',law_url='',location=''):
 return {'scope':scope,'item':item,'status':status,'evidence':clean(evidence)[:3500],'source':source,'source_url':source_url,'document':document,'field_path':field_path,'note':note,'law':law,'law_url':law_url,'location':location}

def analyze(p,docs):
 out=[];deadline=dt(p.get('submission_end') or p.get('deadline'));reqs=[d for d in docs if d['type']=='clarification_request'];answers=[d for d in docs if d['type']=='clarification']
 for rq in reqs:
  rd=mdate(rq,'request') or mdate(rq,'create') or mdate(rq,'publish')
  if not rd or (deadline and wdays(rd,deadline)<3):continue
  later=[]
  for a in answers:
   ad=mdate(a,'publish') or mdate(a,'create')
   if ad and ad>=rd:later.append((ad,a))
  if not later:out.append(finding('CLAR-NO-ANSWER','check','Не найдено разъяснение на своевременный запрос','223-ФЗ, ст. 3.2 ч. 2–3',LAW32,f'Запрос: {rd.isoformat()}; ответа среди доступных документов не найдено.','Контрольный индикатор: требуется убедиться, что в ЕИС действительно отсутствует ответ и запрос был своевременным.',rq));continue
  ad,a=min(later,key=lambda x:x[0]);wd=wdays(rd,ad)
  if wd>3:out.append(finding('CLAR-LATE','violation','Разъяснение опубликовано позднее трёх рабочих дней','223-ФЗ, ст. 3.2 ч. 3',LAW32,f'Запрос: {rd.isoformat()}; разъяснение: {ad.isoformat()}; рабочих дней: {wd}.','Факт срока сформирован по метаданным ЕИС.',a))
 for a in answers:
  m=re.search(r'(?:измен|замен)[^.\n]{0,140}(?:предмет\w*\s+закупк|существенн\w*\s+услов)',a.get('text',''),re.I)
  if m:
   mt=m.group(0);ctx=a['text'][max(0,m.start()-160):m.end()+220]
   out.append(finding('CLAR-CHANGES','risk','Разъяснение может изменять предмет или существенные условия','223-ФЗ, ст. 3.2 ч. 4',LAW32,ctx,'Риск требует сопоставления исходной и изменённой редакций.',a,mt))
 for d in [x for x in docs if x['type']=='cancellation']:
  dec=mdate(d,'decision') or mdate(d,'sign');pub=mdate(d,'publish')
  if dec and pub and dec.date()!=pub.date():out.append(finding('CANCEL-LATE','violation','Решение об отмене размещено не в день принятия','223-ФЗ, ст. 3.2 ч. 6',LAW32,f'Решение: {dec.date()}; размещение: {pub.date()}.','Срок рассчитан по метаданным документа.',d))
  if dec and deadline and dec>deadline:out.append(finding('CANCEL-AFTER-DEADLINE','check','Отмена после окончания подачи заявок требует отдельной правовой проверки','223-ФЗ, ст. 3.2 ч. 5–7',LAW32,f'Окончание подачи: {deadline}; решение: {dec}.','Нужно проверить наличие предусмотренных законом обстоятельств и положение о закупке.',d))
 for d in [x for x in docs if x['type'] in ('protocol','final_protocol')]:
  sign=mdate(d,'sign');pub=mdate(d,'publish')
  if sign and pub and (pub.date()-sign.date()).days>3:out.append(finding('PROTOCOL-LATE','violation','Протокол размещён позднее трёх дней со дня подписания','223-ФЗ, ст. 4 ч. 12',LAW4,f'Подписание: {sign.date()}; размещение: {pub.date()}.','Срок рассчитан по метаданным ЕИС.',d))
  t=d.get('text','')
  if len(re.sub(r'\s+','',t))>=250:
   mandatory=[('дата подписания',r'дата\s+(?:подписан|составлен)|\b\d{2}\.\d{2}\.\d{4}\b'),('количество заявок',r'количеств\w+\s+(?:поданн\w+\s+)?заяв|подано\s+\d+\s+заяв'),('дата/время регистрации',r'регистрац\w+.*(?:дата|время)|дата\s+и\s+время\s+регистрац')]
   if d['type']=='final_protocol':mandatory += [('ранжирование/победитель',r'порядков\w+\s+номер|победител|место\s+\d'),('ценовые предложения',r'ценов\w+\s+предлож|цена\s+(?:договора|заявки)|руб'),('результаты рассмотрения',r'результат\w+\s+рассмотрен|допущ|отклон')]
   for label,pp in mandatory:
    if not re.search(pp,t,re.I|re.S):out.append(finding('PROTOCOL-FIELD','check','В протоколе автоматически не найдено обязательное сведение: '+label,'223-ФЗ, ст. 3.2 ч. 13–14',LAW32,f'Документ «{d["name"]}»: не найдено — {label}.','Отсутствие текста может быть связано с форматом/OCR; до ручного подтверждения отмечается как «требует проверки».',d,location='документ целиком: обязательное сведение автоматически не найдено'))
 # Индикаторы ограничения конкуренции: формулировка -> контекст -> норма. Они помечаются как риск, а не автоматически как установленное нарушение.
 risk_rules=[
  ('DEALER',r'авторизационн\w*\s+письм|статус\s+(?:официального\s+)?дилера','Требование авторизационного письма/дилерского статуса'),
  ('LOCALITY',r'(?:офис|склад|сервисн\w*\s+центр)[^.\n]{0,150}(?:в\s+городе|в\s+регионе|на\s+территории)','Территориальное требование к участнику'),
  ('NO-EQUIV',r'(?:только|исключительно)[^.\n]{0,100}(?:марки|бренда|производител)|без\s+(?:возможности\s+)?(?:поставки\s+)?эквивалент','Указание на конкретный бренд/производителя без явного эквивалента'),
  ('EXPERIENCE',r'опыт\w*\s+(?:исполнения|выполнения|поставки)[^.\n]{0,180}(?:не\s+менее|за\s+последн)','Требование к опыту исполнения аналогичных договоров'),
  ('STAFF',r'наличи\w+\s+в\s+штате[^.\n]{0,160}(?:специалист|работник|персонал)','Требование наличия персонала именно в штате'),
 ]
 for d in docs:
  t=d.get('text','')
  if not t:continue
  for code,pp,title in risk_rules:
   m=re.search(pp,t,re.I|re.S)
   if m:
    mt=m.group(0);ctx=t[max(0,m.start()-180):m.end()+260]
    out.append(finding(code,'risk',title,'223-ФЗ, ст. 3 ч. 1',LAW3,ctx,'Контрольный индикатор ограничения конкуренции. Итоговый вывод зависит от предмета закупки, положения о закупке и объективной необходимости требования.',d,mt))
 return out

def norm_search(s):return re.sub(r'[^0-9a-zа-я]+','',str(s or '').lower().replace('ё','е'))
def _contains_any(text,patterns):
 t=(text or '').lower().replace('ё','е')
 return any(re.search(p,t,re.I|re.S) for p in patterns)

def build_checks(p,detail,card,docs):
 out=[]
 # Реальная проверка вкладок сайта ЕИС: фиксируем HTTP/доступность отдельно от API-источника.
 page_names={'common':'Основная информация','lots':'Лоты','documents':'Документы','changes':'Изменения'}
 for key,label in page_names.items():
  pg=(card.get('pages') or {}).get(key,{})
  ok=pg.get('status')=='ok'
  ev=f"HTTP {pg.get('http_status') or '—'}; конечный URL: {pg.get('final_url') or pg.get('url') or '—'}"
  if pg.get('error'):ev+='; '+pg['error']
  out.append(check('eis-site','Вкладка ЕИС: '+label,'ok' if ok else 'check',ev,'zakupki.gov.ru',pg.get('final_url') or pg.get('url') or p['url'],note='Страница ЕИС реально запрошена сервером и разобрана.' if ok else 'Страница ЕИС не подтверждена; это техническая неполнота проверки.'))
 # Все скалярные поля источника фиксируем, чтобы пользователь видел не «чёрный ящик», а прочитанные значения.
 for path,val in flatten_scalars(detail):out.append(check('card','Поле карточки/API','read',val,'Структурированные данные закупки',p['url'],field_path=path,note='Поле получено и зафиксировано.'))
 if card.get('status')=='ok':
  for label,val in card.get('pairs',[]):out.append(check('card','Поле на странице ЕИС','read',val,'zakupki.gov.ru',p['url'],field_path=label,note='Видимое поле страницы ЕИС прочитано.'))
 else:out.append(check('card','Страница карточки ЕИС','check',card.get('error') or 'Страница недоступна','ЕИС',p['url'],note='Без прямого чтения сайта проверка считается неполной.'))
 key=[('Номер закупки',p.get('reg')),('Наименование закупки',p.get('title')),('Заказчик',p.get('customer')),('Способ закупки',p.get('method')),('НМЦД',p.get('price')),('Дата публикации',p.get('published')),('Начало подачи заявок',p.get('submission_start')),('Окончание подачи заявок',p.get('submission_end'))]
 for label,val in key:out.append(check('card',label,'ok' if clean(val) else 'check',val or 'Не найдено','Карточка закупки',p['url'],note='Ключевое поле найдено.' if clean(val) else 'Ключевое поле не найдено в доступных данных.'))
 pn=p.get('price_num');out.append(check('card','НМЦД распознана как число','ok' if pn is not None and pn>=0 else 'check',str(pn) if pn is not None else p.get('price',''),'Карточка закупки',p['url']))
 st,en=dt(p.get('submission_start')),dt(p.get('submission_end'))
 if st and en:out.append(check('card','Хронология подачи заявок','ok' if en>=st else 'violation',f'{p.get("submission_start")} → {p.get("submission_end")}','Карточка закупки',p['url'],note='Окончание не раньше начала.' if en>=st else 'Окончание подачи раньше начала.',law='223-ФЗ, ст. 4 ч. 9 п. 7 и ч. 10 п. 8',law_url=LAW4))
 else:out.append(check('card','Хронология подачи заявок','check',f'начало: {p.get("submission_start") or "—"}; окончание: {p.get("submission_end") or "—"}','Карточка закупки',p['url'],note='Не удалось проверить обе даты.',law='223-ФЗ, ст. 4 ч. 9 п. 7 и ч. 10 п. 8',law_url=LAW4))

 checked=[d for d in docs if d.get('status')=='checked' and clean(d.get('text'))]
 for d in docs:
  chars=len(clean(d.get('text','')));stt='ok' if d.get('status')=='checked' else 'check';ev=f'Статус: {d.get("status")}; извлечено символов: {chars}; конечный URL: {d.get("final_url") or "—"}'
  if d.get('error'):ev+='; ошибка: '+d['error']
  out.append(check('document','Оригинал документа скачан и прочитан',stt,ev,'Документ закупки',d.get('url') or '',document=d.get('name'),note='Содержимое использовано в правовой проверке.' if stt=='ok' else 'Документ не считается проверенным, пока содержимое не извлечено.'))

 # Контрольный чек-лист по ст. 4 223-ФЗ. Статус «check» = автоматикой не подтверждено, а не готовый вывод о нарушении.
 corpus='\n'.join(d.get('text','') for d in checked)
 site_text=card.get('visible_text','') or ''
 full=(site_text+'\n'+corpus).lower().replace('ё','е')
 rules=[
  ('223-4-9-1','Извещение: способ закупки',[r'способ\w*\s+закуп',r'аукцион|конкурс|запрос\s+(?:котиров|предлож)'],'223-ФЗ, ст. 4 ч. 9 п. 1'),
  ('223-4-9-2','Извещение: сведения о заказчике и контакты',[r'почтов\w+\s+адрес',r'электронн\w+\s+почт|e-?mail',r'телефон'],'223-ФЗ, ст. 4 ч. 9 п. 2'),
  ('223-4-9-3','Извещение: предмет и объём/количество',[r'предмет\w+\s+(?:договора|закупк)',r'количеств|объем'],'223-ФЗ, ст. 4 ч. 9 п. 3'),
  ('223-4-9-4','Извещение: место поставки/выполнения',[r'место\s+(?:поставк|выполнен|оказан)'],'223-ФЗ, ст. 4 ч. 9 п. 4'),
  ('223-4-9-5','Извещение: НМЦД',[r'начальн\w+\s*\(?максимальн|нмцд|начальн\w+\s+цен'],'223-ФЗ, ст. 4 ч. 9 п. 5'),
  ('223-4-9-7','Извещение: начало и окончание подачи заявок',[r'начал\w+\s+подач\w+\s+заяв',r'окончан\w+\s+(?:срока\s+)?подач\w+\s+заяв'],'223-ФЗ, ст. 4 ч. 9 п. 7'),
  ('223-4-10-1','Документация: требования к товару/работе/услуге',[r'техническ\w+\s+характерист',r'требован\w+\s+к\s+(?:качеств|безопасност|товар|работ|услуг)'],'223-ФЗ, ст. 4 ч. 10 п. 1'),
  ('223-4-10-2','Документация: содержание и форма заявки',[r'требован\w+\s+к\s+(?:содержан|форм|оформлен).*заяв',r'форма\s+заявк'],'223-ФЗ, ст. 4 ч. 10 п. 2'),
  ('223-4-10-4','Документация: место, условия и сроки поставки',[r'срок\w+\s+(?:поставк|выполнен|оказан)',r'услови\w+\s+(?:поставк|выполнен|оказан)'],'223-ФЗ, ст. 4 ч. 10 п. 4'),
  ('223-4-10-6','Документация: форма, сроки и порядок оплаты',[r'порядок\s+оплат',r'услови\w+\s+оплат',r'срок\w+\s+оплат'],'223-ФЗ, ст. 4 ч. 10 п. 6'),
  ('223-4-10-7','Документация: обоснование НМЦД',[r'обоснован\w+\s+(?:начальн|нмцд|цен)'],'223-ФЗ, ст. 4 ч. 10 п. 7'),
  ('223-4-10-8','Документация: порядок и сроки подачи заявок',[r'порядок\s+подач\w+\s+заяв',r'дата\s+и\s+время\s+окончан\w+.*подач'],'223-ФЗ, ст. 4 ч. 10 п. 8'),
 ]
 for code,item,patterns,law in rules:
  found=[]
  for pp in patterns:
   m=re.search(pp,full,re.I|re.S)
   if m:found.append(m.group(0))
  ok=len(found)>=min(1,len(patterns))
  out.append(check('223-ФЗ',item,'ok' if ok else 'check','; '.join(found[:4]) if found else 'Автоматически не найдено','Карточка ЕИС + документы',p['url'],field_path=code,note='Требование подтверждено по доступным данным.' if ok else 'Нужно проверить оригинал и положение о закупке; отсутствие совпадения автоматикой само по себе не признаётся нарушением.',law=law,law_url=LAW4))

 # Сверка извещения/карточки и документации (ч. 8 ст. 4).
 if corpus:
  for label,val in [('Номер закупки',p.get('reg')),('Заказчик',p.get('customer')),('Наименование закупки',p.get('title'))]:
   needle=norm_search(val); needle=needle[:80] if len(needle)>80 else needle; found=bool(needle and needle in norm_search(corpus))
   out.append(check('223-ФЗ','Соответствие извещения и документации: '+label,'ok' if found else 'check',val or '—','Карточка ↔ документы',p['url'],note='Совпадение найдено.' if found else 'Автоматически не подтверждено; требуется сопоставление редакций.',law='223-ФЗ, ст. 4 ч. 8',law_url=LAW4))
  if p.get('price_num') is not None:
   variants=[str(int(p['price_num'])),f'{p["price_num"]:.2f}'.replace('.',','),f'{p["price_num"]:,.2f}'.replace(',',' ').replace('.',',')];found=any(norm_search(v) in norm_search(corpus) for v in variants)
   out.append(check('223-ФЗ','Соответствие извещения и документации: НМЦД','ok' if found else 'check',p.get('price'),'Карточка ↔ документы',p['url'],note='НМЦД найдена в документах.' if found else 'НМЦД в таком представлении автоматически не найдена.',law='223-ФЗ, ст. 4 ч. 8',law_url=LAW4))
 else:out.append(check('223-ФЗ','Соответствие извещения и документации','check','Нет документов с извлечённым текстом','Карточка ↔ документы',p['url'],note='Невозможно выполнить ч. 8 ст. 4 без чтения оригиналов.',law='223-ФЗ, ст. 4 ч. 8',law_url=LAW4))
 return out

def annotate(reg,docs,finds,checks):
 for d in docs:
  if not d.get('local'):continue
  relevant=[]
  for f in finds:
   if f.get('document')==d['name'] and f.get('match_text'):relevant.append({'text':f.get('match_text'),'label':f.get('title'),'location':f.get('location')})
  for x in checks:
   if x.get('document')==d['name'] and x.get('status') in ('violation','risk') and x.get('evidence'):relevant.append({'text':compact_seed(x.get('evidence')),'label':x.get('item'),'location':x.get('location')})
  if not relevant:continue
  p=Path(d['local']);folder=FILES/reg/'annotated';folder.mkdir(parents=True,exist_ok=True);dst=folder/(p.stem+'__ПРОВЕРЕНО'+p.suffix)
  try:
   if p.suffix.lower()=='.pdf' and fitz:
    doc=fitz.open(p);hits=0
    for item in relevant:
     seed=compact_seed(item['text'])
     if len(seed)<5:continue
     for pi,pg in enumerate(doc):
      rs=pg.search_for(seed)
      if not rs and len(seed)>45:rs=pg.search_for(seed[:45])
      for rr in rs[:5]:
       pg.add_highlight_annot(rr);hits+=1
       try:pg.add_text_annot(rr.tl,f"Проверка 223-ФЗ: {item['label']}")
       except:pass
      if rs:break
    if hits:doc.save(dst);d['annotated']=str(dst)
    doc.close()
   elif p.suffix.lower()=='.docx' and Document:
    doc=Document(p);hits=0
    for par in doc.paragraphs:
     norm=clean(par.text).lower()
     for item in relevant:
      seed=compact_seed(item['text']).lower()
      if seed and (seed in norm or (len(seed)>45 and seed[:45] in norm)):
       for rr in par.runs:rr.font.highlight_color=WD_COLOR_INDEX.YELLOW
       hits+=1;break
    for table in doc.tables:
     for row in table.rows:
      for cell in row.cells:
       norm=clean(cell.text).lower()
       for item in relevant:
        seed=compact_seed(item['text']).lower()
        if seed and (seed in norm or (len(seed)>45 and seed[:45] in norm)):
         for par in cell.paragraphs:
          for rr in par.runs:rr.font.highlight_color=WD_COLOR_INDEX.YELLOW
         hits+=1;break
    if hits:doc.save(dst);d['annotated']=str(dst)
   elif p.suffix.lower() in ('.xlsx','.xlsm') and openpyxl:
    from openpyxl.comments import Comment
    from openpyxl.styles import PatternFill
    wb=openpyxl.load_workbook(p);hits=0
    for ws in wb.worksheets:
     for row in ws.iter_rows():
      for cell in row:
       val=clean(cell.value).lower()
       for item in relevant:
        seed=compact_seed(item['text']).lower()
        if seed and (seed in val or (len(seed)>45 and seed[:45] in val)):
         cell.fill=PatternFill('solid',fgColor='FFF59D');cell.comment=Comment('Проверка 223-ФЗ: '+item['label'],'ЕИС 223-ФЗ Аудитор');hits+=1;break
    if hits:wb.save(dst);d['annotated']=str(dst)
  except Exception:pass

def save(p,docs,finds,checks,card_payload):
 c=conn();c.execute('''INSERT INTO purchases(reg_number,title,customer,org_form,method,price,published_at,deadline,url,source,last_seen,clarifications,protocols,cancellations,violations,risks,submission_start,submission_end,price_num,card_fields_json,card_source_status) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(reg_number) DO UPDATE SET title=excluded.title,customer=excluded.customer,org_form=excluded.org_form,method=excluded.method,price=excluded.price,published_at=excluded.published_at,deadline=excluded.deadline,url=excluded.url,source=excluded.source,last_seen=excluded.last_seen,clarifications=excluded.clarifications,protocols=excluded.protocols,cancellations=excluded.cancellations,violations=excluded.violations,risks=excluded.risks,submission_start=excluded.submission_start,submission_end=excluded.submission_end,price_num=excluded.price_num,card_fields_json=excluded.card_fields_json,card_source_status=excluded.card_source_status''',(p['reg'],p['title'],p['customer'],org_form(p['customer']),p['method'],p['price'],p['published'],p.get('submission_end') or p.get('deadline'),p['url'],'eis+gosplan',now(),sum(d['type'] in ('clarification','clarification_request') for d in docs),sum(d['type'] in ('protocol','final_protocol') for d in docs),sum(d['type']=='cancellation' for d in docs),sum(f['kind']=='violation' for f in finds)+sum(x['status']=='violation' for x in checks),sum(f['kind']=='risk' for f in finds),p.get('submission_start'),p.get('submission_end'),p.get('price_num'),json.dumps(card_payload,ensure_ascii=False,default=str)[:1000000],card_payload.get('status')))
 c.execute('DELETE FROM docs WHERE reg_number=?',(p['reg'],));c.execute('DELETE FROM findings WHERE reg_number=?',(p['reg'],));c.execute('DELETE FROM checks WHERE reg_number=?',(p['reg'],))
 for d in docs:c.execute('INSERT INTO docs(reg_number,name,url,doc_type,published_at,local_path,annotated_path,status,error,metadata_json,text_chars,final_url,checked_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',(p['reg'],d['name'],d['url'],d['type'],d['published'],d['local'],d['annotated'],d['status'],d['error'],json.dumps(d['meta'],ensure_ascii=False,default=str),len(clean(d.get('text',''))),d.get('final_url'),now()))
 for f in finds:c.execute('INSERT INTO findings(reg_number,code,kind,title,law,law_url,evidence,note,document,source_url,location,match_text) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',(p['reg'],f['code'],f['kind'],f['title'],f['law'],f['law_url'],f['evidence'],f['note'],f['document'],f['source_url'],f.get('location'),f.get('match_text')))
 for x in checks:c.execute('INSERT INTO checks(reg_number,scope,item,status,evidence,source,source_url,document,field_path,note,law,law_url,location) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',(p['reg'],x['scope'],x['item'],x['status'],x['evidence'],x['source'],x['source_url'],x['document'],x['field_path'],x['note'],x.get('law'),x.get('law_url'),x.get('location')))
 c.execute('UPDATE purchases SET documents_found=?,documents_checked=?,checks_ok=?,checks_attention=? WHERE reg_number=?',(len(docs),sum(d.get('status')=='checked' for d in docs),sum(x.get('status')=='ok' for x in checks),sum(x.get('status') in ('check','violation','risk') for x in checks),p['reg']))
 c.commit();c.close();render_report(p,docs,finds,checks)

def render_report(p,docs,finds,checks):
 rows=''.join(f"<tr><td>{html.escape(d['type'])}</td><td>{html.escape(d['name'])}</td><td>{html.escape(d['status'])}</td><td>{html.escape(d.get('error') or '')}</td><td>{'<a href='+repr(d['url'])+'>источник</a>' if d['url'] else '—'}</td></tr>" for d in docs)
 blocks=''.join(f"<article><h3>{html.escape(f['title'])}</h3><p><a href='{html.escape(f['law_url'])}'>{html.escape(f['law'])}</a> · {html.escape(f['kind'])}</p><pre>{html.escape(f['evidence'])}</pre><p>{html.escape(f['note'])}</p></article>" for f in finds)
 ck=''.join(f"<tr><td>{html.escape(x['scope'])}</td><td>{html.escape(x['item'])}</td><td>{html.escape(x['status'])}</td><td>{html.escape(x['evidence'])}</td></tr>" for x in checks[:1500])
 (REPORTS/(p['reg']+'.html')).write_text(f"<!doctype html><meta charset=utf-8><style>body{{font-family:Arial;max-width:1200px;margin:30px}}article{{border-left:5px solid #b33;padding:12px;margin:12px 0}}pre{{white-space:pre-wrap;background:#f5f5f5;padding:10px}}table{{border-collapse:collapse;width:100%}}td,th{{padding:8px;border-bottom:1px solid #ddd;vertical-align:top}}</style><h1>Закупка {p['reg']}</h1><p>{html.escape(p['customer'])}</p><p>НМЦД: {html.escape(p.get('price') or '—')} · начало: {html.escape(p.get('submission_start') or '—')} · окончание: {html.escape(p.get('submission_end') or '—')}</p><a href='{p['url']}'>Карточка ЕИС</a><h2>Нарушения и риски</h2>{blocks or '<p>По выполненным автоматическим правилам нарушений не выявлено.</p>'}<h2>Документы</h2><table>{rows}</table><h2>Что проверено</h2><table>{ck}</table>",encoding='utf-8')

async def process_purchase(client,r,basic=None):
 try:detail=await jget(client,f'/fz223/purchases/{r}')
 except Exception as e:
  detail=basic or {};print(f'AUDIT detail-fallback reg={r} err={e!r}',flush=True)
 card=await fetch_eis_card(client,r,detail)
 eis=card.get('summary') or {}
 st_api,en_api=submission_dates(detail)
 customer=eis.get('customer') if _looks_like_name(eis.get('customer')) else api_customer(detail)
 title=eis.get('title') or txt(getv(detail,'purchaseName','purchaseObjectInfo','purchaseObjectName','title','subject','name'))
 price=eis.get('price') or api_price(detail);pn=eis.get('price_num') if eis.get('price_num') is not None else price_num(price)
 method=eis.get('method') or txt(getv(detail,'purchaseMethodName','purchaseMethod','method','placingWayName'))
 published=eis.get('published') or txt(getv(detail,'publicationDate','publishDate','createDate','publishedAt'))
 st=eis.get('submission_start') or st_api;en=eis.get('submission_end') or en_api
 common=(card.get('pages') or {}).get('common',{}).get('final_url') or old_eis_url(r)
 p={'reg':r,'title':title,'customer':customer,'method':method,'price':price,'price_num':pn,'published':published,'submission_start':st,'submission_end':en,'deadline':en,'url':common}
 docs=merge_docs(extract_docs(detail),await discover_eis_documents(client,r,card));await download_all_documents(client,r,docs)
 finds=analyze(p,docs);checks=build_checks(p,detail,card,docs);annotate(r,docs,finds,checks)
 payload={'status':card.get('status'),'error':card.get('error'),'pairs':card.get('pairs',[]),'visible_text':card.get('visible_text',''),'api_fields':flatten_scalars(detail),'pages':{k:{z:v.get(z) for z in ('url','final_url','status','http_status','error','engine')} for k,v in (card.get('pages') or {}).items()},'discovery':card.get('discovery'),'summary':eis}
 save(p,docs,finds,checks,payload)
 read=sum(d.get('status')=='checked' for d in docs);errs=sum(d.get('status')!='checked' for d in docs)
 print(f'AUDIT reg={r} card={card.get("status")} customer={customer!r} nmcd={pn!r} docs={len(docs)} read={read} unread={errs} findings={len(finds)} checks={len(checks)} url={common}',flush=True)
 if card.get('error'):print(f'AUDIT card-error reg={r} {card.get("error")}',flush=True)
 for d in docs:
  if d.get('status')!='checked':print(f'AUDIT doc-error reg={r} name={d.get("name")!r} status={d.get("status")} err={d.get("error")!r} url={d.get("url")}',flush=True)
 return {'reg':r,'card':card.get('status'),'documents':len(docs),'checked_documents':read,'unread_documents':errs,'findings':len(finds),'checks':len(checks),'customer':customer,'nmcd':pn,'url':common}

def _payload_items(payload):
 if isinstance(payload,list):return payload
 if not isinstance(payload,dict):return []
 for key in ('items','data','content','results','purchases'):
  v=payload.get(key)
  if isinstance(v,list):return v
  if isinstance(v,dict):
   for kk in ('items','data','content','results'):
    if isinstance(v.get(kk),list):return v[kk]
 return []

def _repair_regs(limit=20):
 try:
  c=conn();rows=c.execute("SELECT reg_number FROM purchases WHERE price_num IS NULL OR customer IS NULL OR customer='' OR customer GLOB '[0-9]*' OR card_source_status IS NULL OR card_source_status!='ok' OR documents_checked<documents_found ORDER BY last_seen DESC LIMIT ?",(limit,)).fetchall();c.close();return [r['reg_number'] for r in rows]
 except:return []

async def cycle():
 if lock.locked():return
 async with lock:
  state.update(running=True,last_start=now(),last_error=None,current_reg=None)
  results=[];errors=[]
  try:
   async with httpx.AsyncClient(timeout=httpx.Timeout(35,connect=12),follow_redirects=True,http2=False) as cl:
    fresh=[];back=[]
    try:fresh=_payload_items(await jget(cl,'/fz223/purchases',{'limit':PAGE_SIZE,'skip':0}))
    except Exception as e:errors.append('fresh '+repr(e));print('AUDIT source-error fresh',repr(e),flush=True)
    try:
     back=_payload_items(await jget(cl,'/fz223/purchases',{'limit':PAGE_SIZE,'skip':state['backfill_skip']}));state['backfill_skip']+=PAGE_SIZE
    except Exception as e:errors.append('backfill '+repr(e));print('AUDIT source-error backfill',repr(e),flush=True)
    todo=[];seen=set()
    # Сначала чиним уже показанные закупки без НМЦД/названия/полной проверки.
    for r in _repair_regs(DETAIL_LIMIT):
     if r and r not in seen:seen.add(r);todo.append((r,None))
    for x in fresh+back:
     r=regnum(x)
     if r and r not in seen:seen.add(r);todo.append((r,x))
    for r,basic in todo[:DETAIL_LIMIT]:
     state['current_reg']=r
     try:results.append(await process_purchase(cl,r,basic))
     except Exception as e:
      errors.append(f'{r}: {e!r}');print(f'AUDIT purchase-failed reg={r} err={e!r}',flush=True)
    state['last_result']={'latest':len(fresh),'backfill':len(back),'checked':len(results),'documents':sum(x['documents'] for x in results),'documents_read':sum(x['checked_documents'] for x in results),'documents_unread':sum(x.get('unread_documents',0) for x in results),'checks':sum(x['checks'] for x in results),'errors':errors[:20]};state['last_finish']=now();state['last_error']=' | '.join(errors[:5]) if errors and not results else None
    print('AUDIT cycle',json.dumps(state['last_result'],ensure_ascii=False),flush=True)
  except Exception as e:state['last_error']=repr(e);state['last_finish']=now();print('AUDIT cycle-fatal',repr(e),flush=True)
  finally:state['running']=False;state['current_reg']=None
async def loop():
 await asyncio.sleep(2)
 while True:await cycle();await asyncio.sleep(SCAN_INTERVAL)
@app.on_event('startup')
async def start():
 print(f'AUDIT startup version={APP_VERSION} data={DATA}',flush=True);asyncio.create_task(loop())
@app.get('/health')
async def health():return {'status':'ok','version':APP_VERSION,'features':['auto-scan','eis-card-audit','eis-document-download','document-text-analysis','cross-checks','org-form-filter','nmcd-filter','checked-items','marked-documents','current-eis-notice223','browser-tls-fallback']}
@app.get('/api/status')
async def status():return state
@app.post('/api/scan-now')
async def scan_now():
 if lock.locked():return {'started':False,'message':'Уже выполняется'}
 asyncio.create_task(cycle());return {'started':True}
@app.post('/api/purchases/{reg}/recheck')
async def recheck(reg:str):
 if not re.fullmatch(r'\d{10,}',reg):raise HTTPException(400,'Некорректный номер')
 if lock.locked():raise HTTPException(409,'Сейчас выполняется другая проверка')
 async with lock:
  state.update(running=True,last_start=now(),last_error=None,current_reg=reg)
  try:
   async with httpx.AsyncClient(timeout=httpx.Timeout(35,connect=12),follow_redirects=True) as cl:res=await process_purchase(cl,reg,None)
   state['last_result']=res;state['last_finish']=now();return res
  except Exception as e:state['last_error']=repr(e);state['last_finish']=now();raise HTTPException(500,str(e))
  finally:state['running']=False;state['current_reg']=None

@app.get('/api/org-forms')
async def forms():
 c=conn();counts={r['org_form']:r['count'] for r in c.execute("SELECT org_form,COUNT(*) count FROM purchases GROUP BY org_form")};c.close()
 names=[x[0] for x in ORG]
 for x in counts:
  if x not in names and x not in ('Другая','Не определена',''):names.append(x)
 return {'items':[{'org_form':x,'count':counts.get(x,0)} for x in names if x not in ('Другая','Не определена','')]}
@app.get('/api/procurements')
@app.get('/api/dashboard')
async def procurements(org_form_:list[str]|None=Query(None,alias='org_form'),result_kind:str|None=None,q:str|None=None,price_band:str|None=None,price_min:float|None=None,price_max:float|None=None,limit:int=200):
 c=conn();where=[];args=[]
 forms=[]
 for v in (org_form_ or []):forms.extend([x.strip() for x in str(v).split(',') if x.strip()])
 if forms:
  where.append('org_form IN ('+','.join('?' for _ in forms)+')');args.extend(forms)
 if result_kind=='violations':where.append('violations>0')
 elif result_kind=='risks':where.append('risks>0')
 elif result_kind=='clarifications':where.append('clarifications>0')
 elif result_kind=='protocols':where.append('protocols>0')
 elif result_kind=='cancellations':where.append('cancellations>0')
 if q:where.append('(reg_number LIKE ? OR customer LIKE ? OR title LIKE ?)');z='%'+q+'%';args += [z,z,z]
 bands={'lt500':(None,500000),'gte500':(500000,None),'500to1000':(500000,1000000),'1to5m':(1000000,5000000),'5to10m':(5000000,10000000),'10to50m':(10000000,50000000),'gte50m':(50000000,None)}
 if price_band in bands:
  lo,hi=bands[price_band]
  if lo is not None:where.append('price_num>=?');args.append(lo)
  if hi is not None:where.append('price_num<?' if price_band=='lt500' else 'price_num<=?');args.append(hi)
 if price_min is not None:where.append('price_num>=?');args.append(price_min)
 if price_max is not None:where.append('price_num<=?');args.append(price_max)
 sql='SELECT * FROM purchases'+((' WHERE '+' AND '.join(where)) if where else '')+' ORDER BY last_seen DESC LIMIT ?';args.append(min(1000,max(1,limit)))
 rows=[dict(r) for r in c.execute(sql,args)];c.close();stats={'purchases':len(rows),'violations':sum(r['violations'] or 0 for r in rows),'risks':sum(r['risks'] or 0 for r in rows),'protocols':sum(r['protocols'] or 0 for r in rows),'clarifications':sum(r['clarifications'] or 0 for r in rows),'cancellations':sum(r['cancellations'] or 0 for r in rows)};return {'stats':stats,'purchases':rows}
@app.get('/api/purchases/{reg}')
async def detail_api(reg:str):
 if not re.fullmatch(r'[0-9A-Za-z._-]{1,120}',reg):raise HTTPException(400,'Некорректный номер')
 c=conn();p=c.execute('SELECT * FROM purchases WHERE reg_number=?',(reg,)).fetchone()
 if not p:c.close();raise HTTPException(404,'Не найдено')
 docs=[dict(r) for r in c.execute('SELECT * FROM docs WHERE reg_number=? ORDER BY id',(reg,))]
 for d in docs:
  if d.get('local_path'):d['local_url']=f"/files/{reg}/original/{Path(d['local_path']).name}"
  if d.get('annotated_path'):d['annotated_url']=f"/files/{reg}/annotated/{Path(d['annotated_path']).name}"
 finds=[dict(r) for r in c.execute('SELECT * FROM findings WHERE reg_number=? ORDER BY CASE kind WHEN "violation" THEN 1 WHEN "risk" THEN 2 ELSE 3 END,id',(reg,))]
 checks=[dict(r) for r in c.execute('SELECT * FROM checks WHERE reg_number=? ORDER BY CASE status WHEN "violation" THEN 1 WHEN "risk" THEN 2 WHEN "check" THEN 3 WHEN "ok" THEN 4 ELSE 5 END,id',(reg,))]
 c.close();summary={'documents_total':len(docs),'documents_checked':sum(d['status']=='checked' for d in docs),'documents_failed':sum(d['status']!='checked' for d in docs),'checks_total':len(checks),'checks_ok':sum(x['status']=='ok' for x in checks),'checks_attention':sum(x['status'] in ('check','violation','risk') for x in checks)}
 law_checks=[x for x in checks if (x.get('law') or '').startswith('223-ФЗ') or x.get('scope')=='223-ФЗ']
 return {'purchase':dict(p),'documents':docs,'findings':finds,'checks':checks,'law_checks':law_checks,'summary':summary,'report_url':f'/reports/{reg}.html' if (REPORTS/(reg+'.html')).exists() else None}
@app.get('/reports/{name}')
async def report(name:str):
 if not re.fullmatch(r'[0-9A-Za-z._-]+\.html',name):raise HTTPException(400)
 p=(REPORTS/name).resolve()
 if p.parent!=REPORTS.resolve() or not p.exists():raise HTTPException(404)
 return FileResponse(p)
@app.get('/files/{reg}/{folder}/{name}')
async def file(reg:str,folder:str,name:str):
 if folder not in ('original','annotated') or Path(name).name!=name:raise HTTPException(400)
 root=(FILES/reg/folder).resolve();p=(root/name).resolve()
 if p.parent!=root or not p.exists():raise HTTPException(404)
 return FileResponse(p)

HTML=r'''<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>ЕИС 223-ФЗ Аудитор</title><style>
:root{--bg:#f5f3ed;--ink:#18201b;--green:#155b42;--red:#a22;--amber:#9a6500;--line:#ddd;--muted:#66736c}*{box-sizing:border-box}body{font-family:system-ui;margin:0;background:var(--bg);color:var(--ink)}header{padding:18px 5%;background:#19211c;color:white}main{padding:26px 5%;max-width:1500px;margin:auto}.hero{display:flex;justify-content:space-between;gap:24px;align-items:end}.hero h1{font:46px Georgia;margin:0}.status{background:white;padding:14px;border:1px solid var(--line);max-width:460px}.stats,.filters{display:flex;gap:10px;flex-wrap:wrap;margin:18px 0}.card{background:white;padding:12px 16px;border:1px solid var(--line);min-width:125px}.card b{font-size:26px;display:block}input,select,button,summary{padding:10px;border:1px solid #bbb;background:white;border-radius:5px}button,summary{cursor:pointer}table{width:100%;border-collapse:collapse;background:white}th,td{padding:10px;border-bottom:1px solid var(--line);text-align:left;font-size:13px;vertical-align:top}.bad{color:var(--red);font-weight:700}.modal{position:fixed;inset:0;background:#0008;display:none;padding:3%;overflow:auto}.panel{background:white;max-width:1260px;margin:auto;padding:22px;border-radius:8px}.finding{border-left:5px solid #b36;padding:10px;margin:10px 0;background:#fafafa}.okbox{border-left:5px solid #286b49;padding:10px;background:#f4faf6}.tabs{display:flex;gap:8px;margin:18px 0;flex-wrap:wrap}.tabbtn.active{background:#19211c;color:white}.tab{display:none}.tab.active{display:block}.badge{padding:3px 7px;border-radius:10px;background:#eee;font-size:12px;white-space:nowrap}.s-violation{background:#fee;color:#900}.s-risk,.s-check{background:#fff3cd;color:#765400}.s-ok{background:#e6f5ec;color:#175d3b}.s-read{background:#eef2f4;color:#46555d}pre{white-space:pre-wrap;background:#f4f4f4;padding:10px}a{color:var(--green)}small{color:var(--muted)}details.multi{position:relative;min-width:210px}details.multi[open] .choices{display:grid}.choices{position:absolute;z-index:20;display:none;grid-template-columns:1fr;background:white;border:1px solid var(--line);padding:8px;max-height:320px;overflow:auto;min-width:300px;box-shadow:0 8px 20px #0002}.choices label{padding:6px;display:flex;gap:8px;align-items:center}.law-note{padding:10px;background:#f7f7f2;border-left:4px solid #5c6f64;margin:10px 0}@media(max-width:720px){.hero{display:block}.hero h1{font-size:34px}table{display:block;overflow:auto}.panel{padding:14px}.modal{padding:1%}.choices{min-width:260px;right:0}}
</style><header><b>223 / ЕИС Аудитор · v1.4</b></header><main><div class="hero"><div><h1>Автоматическая проверка 223-ФЗ</h1><p>Прямое чтение карточки ЕИС · оригиналы документов · контрольный чек-лист 223-ФЗ · точное место спорной формулировки</p></div><div class="status" id="status">Статус…</div></div><div class="stats"><div class="card">Закупок<b id="s1">0</b></div><div class="card">Нарушений<b id="s2">0</b></div><div class="card">Рисков<b id="s3">0</b></div></div><div class="filters"><input id="q" placeholder="Поиск"><details class="multi" id="orgmulti"><summary id="orgsummary">Формы</summary><div class="choices" id="orgchecks"></div></details><select id="price"><option value="">Любая НМЦД</option><option value="lt500">До 500 тыс.</option><option value="gte500">От 500 тыс.</option><option value="500to1000">500 тыс. – 1 млн</option><option value="1to5m">1–5 млн</option><option value="5to10m">5–10 млн</option><option value="10to50m">10–50 млн</option><option value="gte50m">От 50 млн</option></select><select id="kind"><option value="">Все результаты</option><option value="violations">Нарушения</option><option value="risks">Риски</option><option value="cancellations">Отмены</option></select><button onclick="scan()">Сканировать сейчас</button></div><table><thead><tr><th>№ закупки</th><th>Заказчик / предмет</th><th>НМЦД</th><th>Результат проверки</th><th></th></tr></thead><tbody id="rows"></tbody></table></main><div class="modal" id="modal" onclick="if(event.target===this)this.style.display='none'"><div class="panel" id="panel"></div></div><script>
const E=s=>document.querySelector(s),esc=s=>String(s??'').replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));async function A(u,o){let r=await fetch(u,o),d=await r.json();if(!r.ok)throw Error(d.detail||r.status);return d}function money(x){let n=Number(x);return Number.isFinite(n)?new Intl.NumberFormat('ru-RU',{maximumFractionDigits:2}).format(n)+' ₽':'—'}function selectedForms(){return [...document.querySelectorAll('#orgchecks input:checked')].map(x=>x.value)}function updateFormSummary(){let a=selectedForms();E('#orgsummary').textContent=a.length?`Формы: ${a.length}`:'Формы'}async function loadForms(){let d=await A('/api/org-forms');E('#orgchecks').innerHTML=d.items.filter(x=>x.org_form&&x.org_form!=='Другая'&&x.org_form!=='Не определена').map(x=>`<label><input type="checkbox" value="${esc(x.org_form)}"> <span>${esc(x.org_form)} <small>(${x.count})</small></span></label>`).join('');document.querySelectorAll('#orgchecks input').forEach(x=>x.onchange=()=>{updateFormSummary();load()})}async function load(){let p=new URLSearchParams();selectedForms().forEach(v=>p.append('org_form',v));if(E('#kind').value)p.set('result_kind',E('#kind').value);if(E('#price').value)p.set('price_band',E('#price').value);if(E('#q').value)p.set('q',E('#q').value);let d=await A('/api/dashboard?'+p);E('#s1').textContent=d.stats.purchases;E('#s2').textContent=d.stats.violations;E('#s3').textContent=d.stats.risks;E('#rows').innerHTML=d.purchases.map(x=>`<tr><td><b>${esc(x.reg_number)}</b></td><td><b>${esc(x.customer)}</b><br><small>${esc(x.title)}</small></td><td><b>${money(x.price_num)}</b></td><td class="${x.violations?'bad':''}">нарушения: ${x.violations||0}<br>риски: ${x.risks||0}<br><small>${esc(x.card_source_status==='ok'?'карточка проверена':x.card_source_status==='partial'?'проверка неполная':'карточка не проверена')} · документы ${x.documents_checked||0}/${x.documents_found||0}</small></td><td><button onclick="openP('${esc(x.reg_number)}')">Открыть проверку</button></td></tr>`).join('')}function tab(n){document.querySelectorAll('.tab,.tabbtn').forEach(x=>x.classList.remove('active'));E('#t'+n).classList.add('active');E('#b'+n).classList.add('active')}function badge(s){let m={violation:'Нарушение',risk:'Риск',check:'Требует проверки',ok:'ОК',read:'Прочитано'};return `<span class="badge s-${esc(s)}">${m[s]||esc(s)}</span>`}async function recheck(r){let b=E('#recheck');b.disabled=true;b.textContent='Проверка…';try{await A('/api/purchases/'+r+'/recheck',{method:'POST'});await openP(r);await load()}catch(e){alert(e.message)}finally{if(E('#recheck')){E('#recheck').disabled=false;E('#recheck').textContent='Перепроверить ЕИС и оригиналы'}}}function whereLine(x){let z=[];if(x.location)z.push('<b>'+esc(x.location)+'</b>');if(x.document)z.push(esc(x.document));if(x.source&&!x.document)z.push(esc(x.source));if(x.source_url)z.push(`<a target="_blank" href="${esc(x.source_url)}">открыть источник</a>`);return z.join('<br>')||'—'}async function openP(r){let d=await A('/api/purchases/'+r),p=d.purchase,s=d.summary;let res=d.findings.length?d.findings.map(f=>`<div class="finding"><b>${esc(f.title)}</b><p>${badge(f.kind)} · ${f.law_url?`<a target="_blank" href="${esc(f.law_url)}">${esc(f.law)}</a>`:esc(f.law||'')}</p>${f.location?`<p><b>Где в документе:</b> ${esc(f.location)}</p>`:''}<pre>${esc(f.evidence)}</pre><p>${esc(f.note)}</p>${f.source_url?`<a target="_blank" href="${esc(f.source_url)}">Исходный документ</a>`:''}</div>`).join(''):`<div class="okbox"><b>По выполненным автоматическим правилам подтверждённых нарушений не выявлено.</b><p>Отдельно просмотрите пункты со статусом «Требует проверки» во вкладке 223-ФЗ.</p></div>`;let docs=d.documents.map(x=>`<tr><td><b>${esc(x.name)}</b></td><td>${badge(x.status==='checked'?'ok':'check')}<br><small>${esc(x.error||'')}</small></td><td>${x.text_chars||0}</td><td>${x.url?`<a target="_blank" href="${esc(x.url)}">ссылка ЕИС/источника</a>`:'—'} ${x.local_url?`<br><a target="_blank" href="${esc(x.local_url)}">скачанный оригинал</a>`:''} ${x.annotated_url?`<br><a target="_blank" href="${esc(x.annotated_url)}"><b>отмеченная копия</b></a>`:''}</td></tr>`).join('');let checks=d.checks.map(x=>`<tr><td>${badge(x.status)}</td><td>${esc(x.scope)}</td><td><b>${esc(x.item)}</b>${x.field_path?`<br><small>${esc(x.field_path)}</small>`:''}</td><td>${esc(x.evidence)}</td><td>${whereLine(x)}</td><td>${esc(x.note||'')}</td></tr>`).join('');let law=(d.law_checks||[]).map(x=>`<tr><td>${badge(x.status)}</td><td><b>${esc(x.item)}</b>${x.law?`<br><small>${x.law_url?`<a target="_blank" href="${esc(x.law_url)}">${esc(x.law)}</a>`:esc(x.law)}</small>`:''}</td><td>${esc(x.evidence)}</td><td>${whereLine(x)}</td><td>${esc(x.note||'')}</td></tr>`).join('');E('#panel').innerHTML=`<button onclick="E('#modal').style.display='none'">Закрыть</button> <button id="recheck" onclick="recheck('${esc(r)}')">Перепроверить ЕИС и оригиналы</button><h2>${esc(p.reg_number)}</h2><p><b>${esc(p.customer)}</b><br>${esc(p.title)}</p><p>НМЦД: <b>${money(p.price_num)}</b> · начало подачи: <b>${esc(p.submission_start||'—')}</b> · окончание: <b>${esc(p.submission_end||'—')}</b></p><p><a target="_blank" href="${esc(p.url)}">Открыть карточку ЕИС</a> ${d.report_url?`· <a target="_blank" href="${d.report_url}">Отчёт</a>`:''}</p><div class="law-note"><b>Методика контрольного анализа:</b> факт → норма 223-ФЗ → доказательство → место в документе → вывод. Статус «Требует проверки» не считается установленным нарушением без проверки контекста и положения о закупке.</div><div class="tabs"><button id="b1" class="tabbtn active" onclick="tab(1)">Нарушения и риски</button><button id="b2" class="tabbtn" onclick="tab(2)">Оригиналы документов</button><button id="b3" class="tabbtn" onclick="tab(3)">Что прочитано в ЕИС</button><button id="b4" class="tabbtn" onclick="tab(4)">Что проверено по 223-ФЗ</button></div><div id="t1" class="tab active">${res}</div><div id="t2" class="tab"><table><thead><tr><th>Документ</th><th>Статус чтения</th><th>Символов</th><th>Источник / отмеченная копия</th></tr></thead><tbody>${docs}</tbody></table></div><div id="t3" class="tab"><table><thead><tr><th>Статус</th><th>Область</th><th>Что прочитано/проверено</th><th>Данные</th><th>Источник</th><th>Комментарий</th></tr></thead><tbody>${checks}</tbody></table></div><div id="t4" class="tab"><table><thead><tr><th>Статус</th><th>Пункт контроля</th><th>Доказательство</th><th>Где проверено</th><th>Вывод</th></tr></thead><tbody>${law}</tbody></table></div>`;E('#modal').style.display='block'}async function status(){let d=await A('/api/status');E('#status').textContent=d.running?('Идёт проверка ЕИС и документов'+(d.current_reg?' · '+d.current_reg:'')):d.last_error?'Ошибка: '+d.last_error:'Автопроверка включена · версия 1.4'}async function scan(){await A('/api/scan-now',{method:'POST'});status()}E('#kind').onchange=load;E('#price').onchange=load;let t;E('#q').oninput=()=>{clearTimeout(t);t=setTimeout(load,300)};loadForms().then(load);status();setInterval(()=>{status();load()},30000);
</script></html>
'''
@app.get('/',response_class=HTMLResponse)
async def home():return HTML
